"""Excel workbooks as a data source: each sheet a table, read into DuckDB on this server.

A workbook reaches Solviz two ways. An admin uploads it (.xlsx) on the Databases page, or gives a SharePoint or
OneDrive link that Solviz reads through Microsoft Graph, with the app registration email already uses or one of
its own. Either way the file is kept here, and each sheet is loaded into a DuckDB file next to it. Every
question is then asked of that file in DuckDB's SQL, read only, as of any other database.

How a sheet becomes a table:

* the header is the first row that names its columns (a title row or a blank one above it is left out);
* each column's type is read from its cells: whole numbers, decimals, dates, date and time, true or false, or
  text, where a cell such as "-", "n/a" or an Excel error ("#DIV/0!") is no value;
* a column and a table get a name SQL can use ("Order Date" becomes order_date), the header as written kept as
  the column's description;
* empty rows are left out, and so is a totals row ("Total", "Grand total"), which would count every figure twice;
* an empty sheet, or one with no rows under its header, is no table, and the load report says so.

A new upload replaces the workbook; a linked one is read again when it changes in SharePoint or OneDrive (its
version tag, checked at most every QUERYBOT_WORKBOOK_REFRESH_SECONDS, 900 by default). A link that cannot be
read keeps the last copy read, and says when that was.
"""

from __future__ import annotations

import base64
import datetime as dt
import hashlib
import io
import json
import logging
import os
import re
import tempfile
import threading
import time
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

log = logging.getLogger("querybot.workbook")

CATALOG = "workbook"       # the tables are read as workbook.main.<sheet>
SCHEMA = "main"
UPLOAD = "upload"
LINK = "link"
MAX_BYTES = int(os.getenv("QUERYBOT_WORKBOOK_MAX_MB") or 100) * 1024 * 1024
GRAPH = "https://graph.microsoft.com/v1.0"
GRAPH_LOGIN = "https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token"


class WorkbookError(ValueError):
    """A workbook that cannot be read, said so an admin can act on it."""


def store_dir() -> Path:
    """Where workbooks and their DuckDB copies are kept: QUERYBOT_WORKBOOKS_DIR, or workbooks/ beside the app."""
    folder = Path(os.getenv("QUERYBOT_WORKBOOKS_DIR") or "workbooks")
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def refresh_seconds() -> int:
    try:
        return max(60, int(os.getenv("QUERYBOT_WORKBOOK_REFRESH_SECONDS") or 900))
    except ValueError:
        return 900


# ── Reading a workbook ───────────────────────────────────────────────────────

@dataclass
class Column:
    name: str              # as SQL reads it: order_date
    header: str            # as the sheet writes it: Order Date
    type: str              # BIGINT, DOUBLE, DATE, TIMESTAMP, BOOLEAN or VARCHAR


@dataclass
class Table:
    sheet: str
    name: str
    columns: list[Column]
    rows: int
    first_row: int         # the sheet's row number of the header (1-based)
    left_out: list[str] = field(default_factory=list)


@dataclass
class Loaded:
    sha256: str
    file_name: str
    tables: list[Table]
    skipped: list[str]
    loaded_at: str

    def shape(self) -> str:
        """The tables and their columns, as a fingerprint: a changed one asks for discovery again."""
        return hashlib.sha256(json.dumps(
            [(t.name, [(c.name, c.type) for c in t.columns]) for t in self.tables]).encode()).hexdigest()[:16]


# Cells that hold no value: a dash, n/a, an Excel error.
_NO_VALUE = re.compile(r"^\s*(?:-+|—|n/?a|na|null|none|nil|#n/a|#value!|#div/0!|#ref!|#name\?|#num!|#null!|"
                       r"#spill!|#calc!)?\s*$", re.IGNORECASE)
# A totals row: its first written cell says so.
_TOTAL = re.compile(r"^\s*(?:grand\s+)?(?:sub)?totals?\b|^\s*sum\b|^\s*total\s+général|^\s*totaux?\b", re.IGNORECASE)


def sql_name(text: str, fallback: str) -> str:
    """ "Order Date" -> order_date, "% Margin" -> margin, "2025 Sales" -> c_2025_sales; a name SQL reads unquoted."""
    words = re.findall(r"[0-9A-Za-zÀ-ÿ]+", str(text or ""))
    name = "_".join(w.lower() for w in words)
    name = re.sub(r"[^0-9a-z_]", lambda m: _fold(m.group(0)), name)
    name = re.sub(r"_+", "_", name).strip("_")
    if not name:
        return fallback
    return f"c_{name}" if name[0].isdigit() else name[:63]


def _fold(char: str) -> str:
    import unicodedata

    return unicodedata.normalize("NFKD", char).encode("ascii", "ignore").decode() or "_"


def _unique(name: str, taken: set[str]) -> str:
    out, n = name, 2
    while out in taken:
        out, n = f"{name}_{n}", n + 1
    taken.add(out)
    return out


def _written(value: Any) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def _header_row(rows: list[tuple]) -> int | None:
    """The first of the top rows that names columns: every written cell text, and at least as many written as
    the widest of the top rows has, less a few (a title row above names one cell)."""
    top = rows[:30]
    widest = max((sum(_written(v) for v in r) for r in top), default=0)
    if widest == 0:
        return None
    for i, row in enumerate(top):
        written = [v for v in row if _written(v)]
        if written and all(isinstance(v, str) for v in written) and len(written) >= max(1, widest * 0.6) \
                and any(_written(v) for r in rows[i + 1:i + 50] for v in r):
            return i
    return None


def _type_of(values: list[Any]) -> str:
    seen = [v for v in values if _written(v) and not (isinstance(v, str) and _NO_VALUE.match(v))]
    if not seen:
        return "VARCHAR"
    if all(isinstance(v, bool) for v in seen):
        return "BOOLEAN"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in seen):
        return "BIGINT"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in seen):
        return "BIGINT" if all(float(v).is_integer() for v in seen) and all(abs(v) < 2 ** 62 for v in seen) \
            else "DOUBLE"
    if all(isinstance(v, (dt.datetime, dt.date)) for v in seen):
        if all(not isinstance(v, dt.datetime) or v.time() == dt.time() for v in seen):
            return "DATE"
        return "TIMESTAMP"
    return "VARCHAR"


def _cell(value: Any, kind: str) -> Any:
    if not _written(value) or isinstance(value, str) and _NO_VALUE.match(value) and kind != "VARCHAR":
        return None
    if kind == "VARCHAR":
        if isinstance(value, dt.datetime):
            return value.date().isoformat() if value.time() == dt.time() else value.isoformat(sep=" ")
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value).strip()
    if kind == "BIGINT":
        return int(value)
    if kind == "DOUBLE":
        return float(value)
    if kind == "DATE":
        return value.date() if isinstance(value, dt.datetime) else value
    return value


def read_tables(data: bytes) -> tuple[list[tuple[Table, list[list[Any]]]], list[str]]:
    """Each sheet with a header and rows under it, as a table and its rows; the sheets left out, and why."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - an unreadable file is said, never raised as a trace
        raise WorkbookError(f"This file could not be read as an Excel workbook (.xlsx): {exc}") from None
    out: list[tuple[Table, list[list[Any]]]] = []
    skipped: list[str] = []
    names: set[str] = set()
    try:
        for sheet in book.worksheets:
            if getattr(sheet, "sheet_state", "visible") != "visible":
                skipped.append(f'"{sheet.title}": a hidden sheet')
                continue
            rows = [tuple(r) for r in sheet.iter_rows(values_only=True)]
            found = _sheet_table(sheet.title, rows, names)
            if isinstance(found, str):
                skipped.append(f'"{sheet.title}": {found}')
            else:
                out.append(found)
    finally:
        book.close()
    if not out:
        raise WorkbookError("No sheet of this workbook has a header row with rows under it. "
                            + ("; ".join(skipped) if skipped else ""))
    return out, skipped


def _sheet_table(title: str, rows: list[tuple], names: set[str]) -> tuple[Table, list[list[Any]]] | str:
    head = _header_row(rows)
    if head is None:
        return "no header row with rows under it" if any(any(_written(v) for v in r) for r in rows) \
            else "an empty sheet"
    header = rows[head]
    spots = [i for i, v in enumerate(header) if _written(v)]
    first, last = spots[0], spots[-1]
    body = [r for r in rows[head + 1:]]
    width = max([last + 1] + [len(r) for r in body]) if body else last + 1
    # Columns past the last header that hold values are kept, under a name of their own.
    while width > last + 1 and not any(len(r) >= width and _written(r[width - 1]) for r in body):
        width -= 1
    taken: set[str] = set()
    heads: list[tuple[int, str, str]] = []
    for i in range(first, width):
        written = str(header[i]).strip() if i < len(header) and _written(header[i]) else ""
        heads.append((i, _unique(sql_name(written, f"column_{i + 1}"), taken), written or f"Column {i + 1}"))
    left_out: list[str] = []
    data: list[list[Any]] = []
    for n, row in enumerate(body, start=head + 2):
        cells = [row[i] if i < len(row) else None for i, _, _ in heads]
        if not any(_written(v) for v in cells):
            continue
        lead = next((v for v in cells if _written(v)), None)
        if isinstance(lead, str) and _TOTAL.match(lead):
            left_out.append(f"a totals row (row {n})")
            continue
        data.append(cells)
    if not data:
        return "no rows under its header"
    if head > 0:
        left_out.insert(0, f"{head} row{'s' if head != 1 else ''} above the header")
    columns = []
    for j, (_, name, written) in enumerate(heads):
        columns.append(Column(name=name, header=written, type=_type_of([r[j] for r in data])))
    typed = [[_cell(r[j], c.type) for j, c in enumerate(columns)] for r in data]
    table = Table(sheet=title, name=_unique(sql_name(title, "sheet"), names), columns=columns, rows=len(typed),
                  first_row=head + 1, left_out=left_out)
    return table, typed


# ── Kept on this server ──────────────────────────────────────────────────────

def _paths(sha: str) -> tuple[Path, Path, Path]:
    if not re.fullmatch(r"[0-9a-f]{64}", sha or ""):
        raise WorkbookError("No workbook has been uploaded for this connection yet.")
    folder = store_dir()
    return folder / f"{sha}.xlsx", folder / f"{sha}.duckdb", folder / f"{sha}.json"


def save(data: bytes, file_name: str) -> Loaded:
    """Keep a workbook's bytes and load its sheets; the same file twice is loaded once."""
    if len(data) > MAX_BYTES:
        raise WorkbookError(f"The workbook is {len(data) / 1024 / 1024:.0f} MB; the most Solviz reads is "
                            f"{MAX_BYTES // 1024 // 1024} MB (QUERYBOT_WORKBOOK_MAX_MB).")
    if not data.startswith(b"PK") or not zipfile.is_zipfile(io.BytesIO(data)):
        raise WorkbookError("This is not an Excel workbook (.xlsx). Save it from Excel as an .xlsx file "
                            "(an older .xls or a .csv is not read).")
    sha = hashlib.sha256(data).hexdigest()
    book, _, _ = _paths(sha)
    if not book.exists():
        _atomic_write(book, data)
    return load(sha, file_name)


_LOADING = threading.Lock()


def load(sha: str, file_name: str = "") -> Loaded:
    """The workbook's DuckDB copy, made once: each sheet a table in workbook.main."""
    import duckdb

    book, built, report = _paths(sha)
    with _LOADING:
        if built.exists() and report.exists():
            return _report(report)
        if not book.exists():
            raise WorkbookError("The workbook this connection reads is no longer on this server: upload it again.")
        tables, skipped = read_tables(book.read_bytes())
        tmp = built.with_suffix(f".{os.getpid()}.{threading.get_ident()}.tmp")
        tmp.unlink(missing_ok=True)
        con = duckdb.connect(str(tmp))
        try:
            for table, rows in tables:
                _write_table(con, table, rows)
            con.execute("CHECKPOINT")
        finally:
            con.close()
        os.replace(tmp, built)
        loaded = Loaded(sha256=sha, file_name=file_name or book.name, tables=[t for t, _ in tables],
                        skipped=skipped, loaded_at=dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"))
        _atomic_write(report, json.dumps(asdict(loaded), default=str).encode())
        return loaded


def _write_table(con, table: Table, rows: list[list[Any]]) -> None:
    import pandas as pd

    frame = pd.DataFrame(rows, columns=[c.name for c in table.columns], dtype=object)
    con.register("sheet_rows", frame)
    try:
        cast = ", ".join(f'CAST("{c.name}" AS {c.type}) AS "{c.name}"' for c in table.columns)
        con.execute(f'CREATE TABLE "{table.name}" AS SELECT {cast} FROM sheet_rows')
    finally:
        con.unregister("sheet_rows")
    for c in table.columns:
        if c.header and c.header != c.name:
            con.execute(f'COMMENT ON COLUMN "{table.name}"."{c.name}" IS {_literal(c.header)}')
    con.execute(f'COMMENT ON TABLE "{table.name}" IS {_literal("Sheet " + table.sheet)}')


def _literal(text: str) -> str:
    return "'" + text.replace("'", "''") + "'"


def _report(path: Path) -> Loaded:
    raw = json.loads(path.read_text(encoding="utf-8"))
    tables = [Table(**{**t, "columns": [Column(**c) for c in t["columns"]]}) for t in raw["tables"]]
    return Loaded(sha256=raw["sha256"], file_name=raw.get("file_name", ""), tables=tables,
                  skipped=list(raw.get("skipped") or []), loaded_at=raw.get("loaded_at", ""))


def _atomic_write(path: Path, data: bytes) -> None:
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False, prefix=path.name, suffix=".part") as handle:
        handle.write(data)
    os.replace(handle.name, path)


# ── A SharePoint or OneDrive link ────────────────────────────────────────────

_LINKS: dict[str, tuple[float, str]] = {}      # link -> (checked at, workbook sha)
_LINKS_LOCK = threading.Lock()


def share_id(link: str) -> str:
    """The id Microsoft Graph reads a sharing link by: "u!" and the link in unpadded base64url."""
    return "u!" + base64.urlsafe_b64encode(link.strip().encode()).decode().rstrip("=")


def _graph_settings(credentials: dict) -> tuple[str, str, str]:
    """The tenant, app and secret a link is read with: the connection's own, or email's (Settings, Email)."""
    if credentials.get("graph") == "mail":
        from core.mail.settings import load as load_mail_settings

        mail = load_mail_settings()
        if mail is None or not (mail.graph_tenant_id and mail.graph_client_id and mail.graph_client_secret):
            raise WorkbookError("Email is not set up to send through Microsoft 365, so its sign-in cannot read "
                                "the link: give this connection a tenant, an app and a secret of its own.")
        return mail.graph_tenant_id, mail.graph_client_id, mail.graph_client_secret
    tenant, app, secret = (str(credentials.get(k) or "").strip() for k in ("tenant_id", "client_id", "client_secret"))
    if not (tenant and app and secret):
        raise WorkbookError("A SharePoint or OneDrive link needs the Microsoft 365 tenant ID, the app's client ID "
                            "and its client secret (an app registration with Files.Read.All).")
    return tenant, app, secret


def _token(http, tenant: str, app: str, secret: str) -> str:
    response = http.post(GRAPH_LOGIN.format(tenant=tenant), data={
        "client_id": app, "client_secret": secret, "scope": "https://graph.microsoft.com/.default",
        "grant_type": "client_credentials"})
    body = _json(response)
    if response.status_code != 200 or not body.get("access_token"):
        detail = str(body.get("error_description") or body.get("error") or response.status_code).splitlines()[0]
        raise WorkbookError(f"Microsoft refused the app's sign-in: {detail[:200]}")
    return str(body["access_token"])


def _json(response) -> dict:
    try:
        body = response.json()
        return body if isinstance(body, dict) else {}
    except Exception:  # noqa: BLE001 - not JSON: nothing to read from it
        return {}


def _refused(response, what: str) -> WorkbookError:
    error = (_json(response).get("error") or {}) if isinstance(_json(response).get("error"), dict) else {}
    detail = str(error.get("message") or response.status_code)[:200]
    if response.status_code in (401, 403):
        return WorkbookError(f"Microsoft 365 would not let the app {what}: {detail}. The app needs the "
                             "Files.Read.All (or Sites.Read.All) application permission, with admin consent.")
    if response.status_code == 404:
        return WorkbookError(f"No file was found at this link ({detail}). Copy the link again from SharePoint or "
                             "OneDrive (Share, Copy link).")
    return WorkbookError(f"Microsoft 365 did not {what} ({detail}).")


def read_link(credentials: dict, *, http=None, force: bool = False) -> Loaded:
    """The workbook a SharePoint or OneDrive link points at, read again when its version tag has changed (checked
    at most every refresh_seconds()). A link that cannot be read keeps the last copy read."""
    link = str(credentials.get("link") or "").strip()
    if not re.match(r"^https://", link):
        raise WorkbookError("Paste the file's SharePoint or OneDrive link (it starts with https://).")
    with _LINKS_LOCK:
        held = _LINKS.get(link)
    if held and not force and time.time() - held[0] < refresh_seconds():
        return load(held[1])
    kept = held[1] if held else _kept_sha(link)
    try:
        sha = _fetch(link, credentials, http, kept)
    except WorkbookError:
        if kept:
            log.warning("workbook: %s could not be read again; answering from the copy read last", link)
            return load(kept)
        raise
    with _LINKS_LOCK:
        _LINKS[link] = (time.time(), sha)
    return load(sha)


def _link_note(link: str) -> Path:
    return store_dir() / f"link-{hashlib.sha256(link.encode()).hexdigest()[:24]}.json"


def _kept_sha(link: str) -> str:
    try:
        return str(json.loads(_link_note(link).read_text(encoding="utf-8")).get("sha256") or "")
    except (OSError, ValueError):
        return ""


def _fetch(link: str, credentials: dict, http, kept: str) -> str:
    own = http is None
    if own:
        import httpx

        http = httpx.Client(timeout=60, follow_redirects=True)
    try:
        tenant, app, secret = _graph_settings(credentials)
        headers = {"Authorization": f"Bearer {_token(http, tenant, app, secret)}"}
        item = f"{GRAPH}/shares/{share_id(link)}/driveItem"
        meta = http.get(item, params={"$select": "name,eTag,cTag,size,file,lastModifiedDateTime"}, headers=headers)
        if meta.status_code != 200:
            raise _refused(meta, "read the link")
        info = _json(meta)
        name = str(info.get("name") or "workbook.xlsx")
        if not name.lower().endswith((".xlsx", ".xlsm")):
            raise WorkbookError(f'The link points at "{name}", not an Excel workbook (.xlsx).')
        tag = str(info.get("cTag") or info.get("eTag") or "")
        note = _link_note(link)
        try:
            known = json.loads(note.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            known = {}
        if kept and tag and known.get("tag") == tag:
            return kept
        content = http.get(f"{item}/content", headers=headers)
        if content.status_code != 200:
            raise _refused(content, "download the file")
        loaded = save(content.content, name)
        _atomic_write(note, json.dumps({"tag": tag, "sha256": loaded.sha256, "name": name,
                                        "modified": info.get("lastModifiedDateTime")}).encode())
        return loaded.sha256
    except WorkbookError:
        raise
    except Exception as exc:  # noqa: BLE001 - the network: said, the last copy kept
        raise WorkbookError(f"Microsoft 365 could not be reached: {exc}") from None
    finally:
        if own:
            http.close()


# ── Answering from it ────────────────────────────────────────────────────────

def current(credentials: dict, *, http=None) -> Loaded:
    """The workbook a connection answers from now: its upload, or its link's latest version."""
    if credentials.get("source") == LINK:
        return read_link(credentials, http=http)
    return load(str(credentials.get("workbook") or ""), str(credentials.get("file_name") or ""))


def connect(credentials: dict, *, http=None):
    """A read-only DuckDB connection with the workbook's tables as workbook.main.<sheet>."""
    import duckdb

    loaded = current(credentials, http=http)
    _, built, _ = _paths(loaded.sha256)
    con = duckdb.connect()
    con.execute(f"ATTACH '{built.as_posix()}' AS {CATALOG} (READ_ONLY)")
    con.execute(f"USE {CATALOG}")
    # Nothing a query sends reaches another file of this server (read_csv('/etc/...'), COPY ... TO, ATTACH,
    # INSTALL), and the setting cannot be turned back on: the workbook's own tables are all it reads.
    con.execute("SET enable_external_access = false")
    con.execute("SET lock_configuration = true")
    return con


def describe(credentials: dict) -> dict:
    """What the Test button shows: the file, its tables and rows, and what was left out."""
    loaded = current(credentials)
    return {"workbook": loaded.file_name, "read": loaded.loaded_at,
            "tables": ", ".join(f"{t.name} ({t.rows:,} rows)" for t in loaded.tables),
            **({"left out": "; ".join(loaded.skipped)} if loaded.skipped else {})}
