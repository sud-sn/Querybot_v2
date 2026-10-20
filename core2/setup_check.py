"""The setup check: what Learn found, as a workbook the client's own people answer, and their answers back.

Learn reads a database well enough to answer from it, but not the business behind it: that reversed transactions
never count, that a loan term is averaged and not added up, that "scripts" are fills. The people who know are the
client's own, and they will not learn the admin console. So the model goes out as a workbook in their words --
each row what Solviz found, with numbers from the data to judge it by, and a yellow cell to answer -- and comes
back as admin decisions (store.set_core2_override), written exactly as the Data pages write them, so they hold
through the next Learn.

    export(model, ...)  -> bytes of the workbook, every answerable row carrying a hidden id
    read(data, ...)     -> the answers, by row id, from the workbook the client filled in
    review(model, ...)  -> the changes the answers propose, each checked on the data, for an admin to apply
    apply(...)          -> the chosen changes, written as decisions, requests and test questions

Nothing an answer says is applied unseen: review() only proposes, and an admin picks what apply() writes. A
correction in the client's own words, which no field can hold, goes to the Requests queue for an admin.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import logging
import re
import time
from typing import Any

from core2.model.schema import AggExpr, ColumnFilter, SemanticModel

log = logging.getLogger("querybot.core2")

VERSION = "solviz-setup-check-1"
META = "_solviz"
LISTS = "_lists"
ID_HEADER = "Solviz ID (keep)"
HEAD_ROW, FIRST_ROW = 4, 5

SHOWN_AS = {"currency": "Money", "count": "Count", "integer": "Whole number", "number": "Number",
            "percent": "Percent"}
ADD_UP = {"sum": "Total", "avg": "Average"}
KEPT, LEFT_OUT = "Kept as Unknown", "Left out"
YES_NO = ["Yes", "No", "Not sure"]

# Rows that may not belong in totals: a status or flag that says the row was undone, failed or is not real.
_UNDONE = re.compile(r"revers|returned|\bvoid|cancel|reject|fail|declin|\btest\b|delet|duplicat|bounced|"
                     r"written[ _-]?off|\bnsf\b", re.I)
_A_STATUS = re.compile(r"status|state\b|_st$|stat_?cd", re.I)
_CODE = re.compile(r"[A-Za-z0-9_\-]{1,6}")
_NOT_A_STATUS = re.compile(r"reason|type|categor|name|descr|bucket|channel|comment|note", re.I)


class SetupCheckError(ValueError):
    """The workbook cannot be read as this workspace's setup check, said in words."""


# ── names ───────────────────────────────────────────────────────────────────

def _table(model: SemanticModel, key: str) -> str:
    t = model.tables.get(key)
    return (t.business_name or t.name) if t else key.rsplit(".", 1)[-1]


def _column(model: SemanticModel, key: str) -> str:
    c = model.columns.get(key)
    return (c.business_name or c.name) if c else key.rsplit(".", 1)[-1]


def _physical(key: str) -> str:
    return ".".join(key.split(".")[-2:])


def _plural(word: str) -> str:
    w = word.strip()
    if not w or w.endswith("s"):
        return w
    return w[:-1] + "ies" if w.endswith("y") and w[-2:-1] not in "aeiou" else w + "s"


def _dates(model: SemanticModel, table: str) -> list[tuple[str, str]]:
    """A table's dates a figure may be counted by, as (name shown, key): names made unique, in a stable order."""
    out: list[tuple[str, str]] = []
    seen: dict[str, int] = {}
    for role in sorted((r for r in model.date_roles.values() if r.table == table and r.kind != "audit"),
                       key=lambda r: (not r.is_default, r.key)):
        name = (role.name or role.key).replace(",", " ").strip()
        seen[name] = seen.get(name, 0) + 1
        out.append((name if seen[name] == 1 else f"{name} ({seen[name]})", role.key))
    return out


def _borrowed(model: SemanticModel, table: str) -> str:
    """The date a table with none of its own is counted by: its document's ("Order date (from Order)"), as answers
    count an order line by its order's date; "" when there is no one document to take it from."""
    from core2.resolve.resolver import _documents

    if table not in model.tables or any(r.table == table and r.kind in ("event", "snapshot")
                                        for r in model.date_roles.values()):
        return ""
    documents = _documents(model, table)
    role = model.date_roles.get(model.tables[documents[0]].default_date or "") if len(documents) == 1 else None
    return f"{role.name or role.key} (from {_table(model, documents[0])})" if role else ""


def _date_shown(model: SemanticModel, table: str) -> str:
    """The date a table is counted by, as the workbook says it."""
    own = next((n for n, k in _dates(model, table) if k == model.tables[table].default_date), "")
    return own or _borrowed(model, table) or "None found"


def _words(text: Any) -> list[str]:
    return list(dict.fromkeys(w.strip() for w in re.split(r"[,;\n]", str(text or "")) if w.strip()))[:12]


def _id(*parts: Any) -> str:
    return json.dumps(list(parts), default=str)


def _filter_id(f: ColumnFilter) -> str:
    return hashlib.sha1(json.dumps(f.model_dump(), sort_keys=True, default=str).encode()).hexdigest()[:12]


def _typed(model: SemanticModel, column_key: str, value: Any) -> Any:
    """A value as its column holds it: a flag as true or false, a number as a number."""
    column = model.columns[column_key]
    text = str(value).strip()
    try:
        if column.data_type == "boolean":
            return text.lower() in ("true", "1", "yes", "y", "t")
        if column.data_type == "integer":
            return int(float(text))
        if column.data_type in ("decimal", "float"):
            return float(text)
    except ValueError:
        pass
    return text


# ── what goes in the workbook ───────────────────────────────────────────────

def leave_out_rows(model: SemanticModel) -> list[dict[str, Any]]:
    """Rows a table may not want in its totals: what it leaves out now, and statuses that read as undone."""
    from core2.model.links import condition_text

    out: list[dict[str, Any]] = []
    for t in model.tables.values():
        if t.kind in ("calendar", "dimension") or t.hidden or not (t.default_date or _borrowed(model, t.key)):
            continue                                   # a lookup's rows are in no total
        now = {(f.column, str(v).lower()) for f in t.default_filters for v in f.values}
        for f in t.default_filters:
            out.append({"id": _id("leftout", t.key, _filter_id(f)), "table": t.key, "words": condition_text(model, f),
                        "rows": None, "share": None, "now": LEFT_OUT, "technical": f"{_physical(f.column)} {f.op} "
                        f"{', '.join(map(str, f.values))}"})
        for c in model.table_columns(t.key):
            p = c.profile
            if c.hidden or c.sensitivity != "none" or c.personal != "none" or c.record_number \
                    or not c.values_allowed or p is None or not p.top or _NOT_A_STATUS.search(c.name):
                continue
            yes_no = c.data_type == "boolean" or c.role == "flag"
            # A status held as short codes says nothing in its values: each code is asked about, less urgently.
            coded = c.data_type == "text" and bool(_A_STATUS.search(c.name)) and 2 <= len(p.top) <= 12 \
                and all(t.value is not None and _CODE.fullmatch(str(t.value)) for t in p.top)
            for top in p.top:
                value, text = top.value, str(top.value).strip()
                if value is None or (c.key, text.lower()) in now:
                    continue
                flag = yes_no and text.lower() in ("true", "1", "y", "yes") and bool(_UNDONE.search(c.name))
                undone = not yes_no and bool(_UNDONE.search(text) or _UNDONE.search(c.value_names.get(text, "")))
                if not (flag or undone or coded):
                    continue
                label = re.sub(r"^is |\s*flag$", "", _column(model, c.key).lower()).strip()
                things = _plural(_table(model, t.key)).capitalize()
                named = c.value_names.get(text)
                out.append({"id": _id("leave", t.key, c.key, text), "table": t.key,
                            "words": f"{things} marked {label}" if flag else
                            f"{things} with {label} \"{text}\"" + (f" ({named})" if named else ""),
                            "rows": top.count, "share": top.count / t.row_count if t.row_count else None,
                            "now": "Counted in totals", "technical": f"{_physical(c.key)} = {text}",
                            "optional": not (flag or undone)})
    return sorted(out, key=lambda r: (r["now"] == LEFT_OUT, r.get("optional", False)))


def figures_month(model: SemanticModel, today: dt.date) -> tuple[dt.date, dt.date]:
    """The month the workbook shows figures for: the last whole month of the biggest table's data (a table of
    plans or targets may run on into months still to come), never one after today."""
    roles = [r for r in model.date_roles.values() if r.is_default and r.last and r.kind != "audit"
             and r.table in model.tables and model.tables[r.table].kind in ("fact", "snapshot")]
    biggest = max(roles, key=lambda r: (model.tables[r.table].row_count, r.key), default=None)
    latest = min(biggest.last if biggest else today, today)
    if (latest + dt.timedelta(days=1)).month == latest.month:
        latest = latest.replace(day=1) - dt.timedelta(days=1)
    return latest.replace(day=1), latest


def _figure(model: SemanticModel, slug: str, warehouse: Any, start: dt.date, end: dt.date, today: dt.date) -> Any:
    """A metric over one month, as an answer would work it out (its own date, its table's leave-out rules)."""
    from core2.compile.compiler import compile_query
    from core2.plan.ir import Plan
    from core2.resolve.resolver import Context, resolve

    from core2.resolve.resolver import ResolveError

    plan = {"kind": "query", "measures": [slug], "time": {"window": {
        "kind": "between", "start": start.isoformat(), "end": end.isoformat()}}}
    try:
        logical, in_all = resolve(Plan.model_validate(plan), model, Context(today=today)), False
    except ResolveError:
        # Members of a list in a month are counted by what they did or by a date of their own: an answer asks
        # which; the workbook shows them all.
        logical, in_all = resolve(Plan.model_validate({**plan, "time": {}}), model, Context(today=today)), True
    compiled = compile_query(logical, model, warehouse.dialect)
    result = warehouse.query(compiled.sql, max_rows=2)
    name = next(c.name for c in compiled.columns if c.role == "measure")
    value = dict(zip(result.columns, result.rows[0])).get(name) if result.rows else None
    return (value, "in all") if in_all else (value, "")


def _shown_value(value: Any, fmt: str) -> str:
    if value is None:
        return "—"
    try:
        x = float(value)
    except (TypeError, ValueError):
        return str(value)
    if fmt == "currency":
        return f"${x:,.0f}"
    if fmt in ("count", "integer"):
        return f"{x:,.0f}"
    if fmt == "percent":
        return f"{x:,.1f}%"
    return f"{x:,.2f}"


def _number(text: Any) -> float | None:
    cleaned = re.sub(r"[^0-9.\-]", "", str(text or ""))
    try:
        return float(cleaned) if cleaned not in ("", "-", ".") else None
    except ValueError:
        return None


_SUMMED_LIKE_A_RATE = re.compile(r"\b(rate|score|term|percent(age)?|ratio|age|tenure|price)s?\b"
                                 r"|\bdays? (past|late|overdue|outstanding|open|to)\b", re.I)


def metric_rows(model: SemanticModel) -> list[dict[str, Any]]:
    """Every metric readers can ask for, with what it works out and anything about it that looks off."""
    from core2.plan.catalog import definition

    out = []
    for m in model.measures.values():
        if m.hidden or re.search(r"\brows?$", m.business_name, re.I):
            continue                                   # row counts of technical tables
        expr, notes = m.expr, []
        column = _column(model, expr.column) if isinstance(expr, AggExpr) and expr.column else ""
        adds = ADD_UP.get(expr.agg) if isinstance(expr, AggExpr) and expr.column else None
        if adds == "Total" and _SUMMED_LIKE_A_RATE.search(f"{column} {m.business_name}"):
            notes.append("This adds the values up across rows. Teams usually want the average: should it be averaged?")
        if m.format == "currency" and re.search(r"\b(count|score|number)\b", f"{m.business_name} {column}", re.I):
            notes.append("Shown as money, but it looks like a count or a score.")
        if m.format != "currency" and re.search(r"balance|amount|revenue|sales|cost|price", m.business_name, re.I):
            notes.append("Looks like money, but is not shown as money.")
        dates = _dates(model, m.table)
        date_key = m.default_date or (model.tables[m.table].default_date if m.table in model.tables else None)
        borrowed = "" if dates else _borrowed(model, m.table)
        if not dates and not borrowed and m.table in model.tables and model.tables[m.table].kind in ("fact", "snapshot"):
            notes.append("No date found to count it by: which period does each row belong to?")
        try:
            works_out = definition(model, expr)
        except Exception:  # noqa: BLE001 - a formula the words cannot follow is said by its name
            works_out = m.description or m.business_name
        out.append({"id": _id("metric", m.key), "key": m.key, "slug": m.slug, "name": m.business_name,
                    "works_out": works_out, "adds_up": adds, "format": m.format, "dates": dates,
                    "date": next((n for n, k in dates if k == date_key), dates[0][0] if dates else borrowed),
                    "table": _table(model, m.table),
                    "notes": notes, "synonyms": list(dict.fromkeys(w for ws in m.synonyms.values() for w in ws))})
    # Two metrics of one name (a quantity on the lines and on the fills) are told apart by their table.
    seen: dict[str, int] = {}
    for row in out:
        seen[row["name"].casefold()] = seen.get(row["name"].casefold(), 0) + 1
    for row in out:
        row["shown"] = f"{row['name']} ({row['table']})" if seen[row["name"].casefold()] > 1 else row["name"]
    return sorted(out, key=lambda r: (not r["notes"], r["shown"].casefold()))


def code_rows(model: SemanticModel) -> list[dict[str, Any]]:
    """Short codes readers will see in answers: what each means, from its field's own names when it has them."""
    from core2.model.knowledge import namable

    keys = {k for j in model.joins.values() for k in (*j.from_columns, *j.to_columns)}
    keys |= {e.code_column for e in model.entities.values() if e.code_column}
    out = []
    for c in model.columns.values():
        t = model.tables.get(c.table)
        if t is None or t.kind == "calendar" or t.hidden or c.hidden or c.personal != "none" or c.record_number \
                or c.key in keys or c.role in ("key", "foreign_key", "identifier", "flag", "label"):
            continue                                   # a key is named by its own table; a yes/no says itself
        values = namable(c)
        if {v["code"].upper() for v in values} <= {"Y", "N", "T", "F", "0", "1", "TRUE", "FALSE", "YES", "NO"} \
                or (len(values) > 3 and all(v["rows"] == 1 for v in values)) \
                or not all(v["code"].isupper() or v["code"].isdigit() or len(v["code"]) <= 3 for v in values):
            continue                                   # one per member is that member's own code; words say themselves
        if len(values) < 2 or not all(_CODE.fullmatch(v["code"]) for v in values) \
                or all(v["code"].isdigit() and len(v["code"]) == 4 for v in values):
            continue                                   # long names say themselves; four digits are years
        for v in values:
            out.append({"id": _id("code", c.key, v["code"]), "where": f"{_table(model, c.table)}: {_column(model, c.key)}",
                        "code": v["code"], "rows": v["rows"], "named": v["name"],
                        "technical": f"{_physical(c.key)} = {v['code']}"})
    return out


def suggestion_rows(model: SemanticModel, account_id: str, most: int = 40) -> list[dict[str, Any]]:
    """New metrics waiting for an admin (Learn's proposals, readers' own): the client says whether they want them."""
    import store
    from core2.model.schema import Measure
    from core2.plan.catalog import definition

    out = []
    for r in store.list_core2_requests(account_id, status="waiting", limit=500):
        if r.get("kind") != "new_metric":
            continue
        made = r.get("definition") or {}
        works_out, fmt, date = str(r.get("note") or ""), str(made.get("format") or ""), ""
        try:
            measure = Measure.model_validate(made)
            if not works_out:
                works_out = definition(model, measure.expr)
            role = model.date_roles.get(measure.default_date or model.tables[measure.table].default_date or "")
            date = (role.name or role.key) if role else ""
        except Exception:  # noqa: BLE001 - a definition this model cannot word is said in the note's words
            pass
        months = (made.get("checked") or {}).get("months") or []
        out.append({"id": int(r["id"]), "name": str(r.get("target_name") or ""), "works_out": works_out,
                    "format": fmt, "date": date,
                    "figure": f"{months[-1].get('shown')} ({months[-1].get('period')})" if months else ""})
    return out[:most]


# ── the workbook ────────────────────────────────────────────────────────────

def export(model: SemanticModel, *, workspace: str, account_id: str, today: dt.date | None = None,
           warehouse: Any = None, suggestions: list[dict[str, Any]] | None = None, budget_s: float = 25.0) -> bytes:
    """The workbook a client fills in: what Solviz found, in their words, with numbers to judge it by."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    today = today or dt.date.today()
    month_start, month_end = figures_month(model, today)
    month = f"{month_end:%b %Y}"
    font, ink, muted = "Arial", "1F2A33", "7A8791"
    head = PatternFill("solid", fgColor="0F5F63")
    answer = PatternFill("solid", fgColor="FFF2B3")
    fills = {"Please review": PatternFill("solid", fgColor="FBE3D0"), "Optional": PatternFill("solid", fgColor="EEF2F4"),
             "Suggested": PatternFill("solid", fgColor="DCEFE6")}
    thin = Side(style="thin", color="D5DDE2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title: str, headers: list[str], widths: list[int], who: str):
        ws = wb.create_sheet(title)
        ws.sheet_view.showGridLines = False
        ws["A1"], ws["A2"] = title, who
        ws["A1"].font = Font(name=font, size=14, bold=True, color=ink)
        ws["A2"].font = Font(name=font, size=10, italic=True, color=muted)
        for i, (h, w) in enumerate(zip(headers, widths), start=1):
            cell = ws.cell(row=HEAD_ROW, column=i, value=h)
            cell.font, cell.fill, cell.border = Font(name=font, size=10, bold=True, color="FFFFFF"), head, border
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        if headers[-1] == ID_HEADER:
            ws.column_dimensions[get_column_letter(len(headers))].hidden = True
        ws.row_dimensions[HEAD_ROW].height = 32
        ws.freeze_panes = f"A{FIRST_ROW}"
        return ws

    def put(ws, row: int, values: list[Any], answers: tuple[int, ...] = (), quiet: tuple[int, ...] = (),
            priority: str | None = None):
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=row, column=i, value=v)
            if isinstance(v, str) and v.startswith("="):
                cell.data_type = "s"          # a name or code from the data is text in the workbook, never a formula
            cell.font = Font(name=font, size=10, color=muted if i in quiet else ink)
            cell.alignment, cell.border = wrap, border
            if i in answers:
                cell.fill = answer
        if priority:
            first = ws.cell(row=row, column=1)
            first.fill = fills[priority]
            first.font = Font(name=font, size=10, bold=priority != "Optional", color=ink)

    def choices(ws, ref: str, options: list[str] | str, *, strict: bool = True):
        formula = options if isinstance(options, str) else '"' + ",".join(o.replace('"', "") for o in options) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=strict)
        dv.error, dv.errorTitle = "Pick one of the choices in the list.", "Choose from the list"
        ws.add_data_validation(dv)
        dv.add(ref)

    # Start here
    ws = wb.create_sheet("Start here")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 26, 90
    lines: list[tuple[str, Any]] = [
        ("title", f"Solviz setup check: {workspace}"),
        ("text", "Solviz has read your database and worked out what your tables, links and figures mean. Please "
                 "check what it found, so every answer it gives uses your business's own definitions."),
        ("text", "It takes about 1 to 2 hours. Most rows only need one click."), ("gap", ""),
        ("head", "How to answer"),
        ("text", "Each row shows what Solviz found, with numbers from your data so you can judge it. In the yellow "
                 "cells pick from the list, or write what is right. Leave a row blank to keep what Solviz found."),
        ("text", "Nothing changes until your Solviz admin uploads this file and approves the changes. Please do not "
                 "add, delete or move tabs."), ("gap", ""),
        ("head", "Colours"),
        ("swatch", (answer, "Yellow cells: your answers.")),
        ("swatch", (fills["Please review"], "Please review: rows where your answer matters most.")),
        ("swatch", (fills["Suggested"], "Suggested: something Solviz thinks you may want added.")),
        ("swatch", (fills["Optional"], "Optional: look if you have time.")),
        ("text", "Grey text: technical names, for your IT team. You do not need to read them."), ("gap", ""),
        ("head", "Who fills which tab"),
        ("pair", ("1 Rows to leave out", "Business owner: rows that should not count in totals.")),
        ("pair", ("2 Metrics", "Business owner: what each figure means and how it is worked out.")),
        ("pair", ("3 Tables and dates", "Business owner or IT: what each table holds and which date it is counted by.")),
        ("pair", ("4 Links", "IT team: how the tables link to each other.")),
        ("pair", ("5 Codes", "Business owner or IT: what short codes mean.")),
        ("pair", ("6 Your words", "Business owner: words your team uses that mean something specific.")),
        ("pair", ("7 Test questions", "Business owner: real questions with the answer from your own reports. Solviz is "
                                      "checked against them before your team uses it.")),
    ]
    for r, (kind, value) in enumerate(lines, start=1):
        if kind == "title":
            ws.cell(row=r, column=1, value=value).font = Font(name=font, size=16, bold=True, color=ink)
        elif kind == "head":
            ws.cell(row=r, column=1, value=value).font = Font(name=font, size=12, bold=True, color="0F5F63")
        elif kind == "text":
            c = ws.cell(row=r, column=1, value=value)
            c.font, c.alignment = Font(name=font, size=10, color=ink), wrap
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            ws.row_dimensions[r].height = 30 if len(value) > 110 else 16
        elif kind == "swatch":
            ws.cell(row=r, column=1).fill = value[0]
            ws.cell(row=r, column=2, value=value[1]).font = Font(name=font, size=10, color=ink)
        elif kind == "pair":
            ws.cell(row=r, column=1, value=value[0]).font = Font(name=font, size=10, bold=True, color=ink)
            b = ws.cell(row=r, column=2, value=value[1])
            b.font, b.alignment = Font(name=font, size=10, color=ink), wrap

    # 1 Rows to leave out
    ws = sheet("1 Rows to leave out", ["Priority", "Table", "These rows", "How many", "Share of the table",
                                       "Solviz now", "Leave them out of totals?", "Your note", "Technical (for IT)",
                                       ID_HEADER],
               [15, 18, 46, 12, 12, 16, 16, 34, 34, 10],
               "Business owner. Rows that may not belong in totals: reversed, returned, cancelled or test rows.")
    r = FIRST_ROW
    for row in leave_out_rows(model):
        left_out = row["now"] == LEFT_OUT
        priority = "Optional" if left_out or row.get("optional") else "Please review"
        put(ws, r, [priority, _table(model, row["table"]), row["words"],
                    f"{row['rows']:,}" if row["rows"] is not None else "",
                    f"{100 * row['share']:.1f}%" if row["share"] is not None else "", row["now"],
                    "Yes" if left_out else None, None, row["technical"], row["id"]],
            answers=(7, 8), quiet=(9,), priority=priority)
        r += 1
    for _ in range(3):
        put(ws, r, ["", "", "(Write any other rows that should not count)", "", "", "", None, None, "", ""],
            answers=(3, 7, 8))
        r += 1
    choices(ws, f"G{FIRST_ROW}:G{r}", YES_NO)

    # 2 Metrics
    started = time.monotonic()
    ws = sheet("2 Metrics", ["Priority", "Metric", "What Solviz works out", "Add up as", "Shown as", "Counted by date",
                             f"Solviz's figure, {month}", "Solviz note", "Right?", "What it should be",
                             f"Your own figure, {month} (optional)", "Other names your team uses", "Hide it?",
                             ID_HEADER],
               [15, 22, 36, 11, 13, 18, 16, 38, 10, 34, 18, 26, 9, 10],
               "Business owner. What each figure means. Check the Solviz notes first; change a yellow cell only "
               "where it is wrong.")
    r = FIRST_ROW
    for row in metric_rows(model):
        figure, said = None, ""
        if warehouse is not None and time.monotonic() - started < budget_s:
            try:
                figure, said = _figure(model, row["slug"], warehouse, month_start, month_end, today)
            except Exception as exc:  # noqa: BLE001 - a figure that cannot be worked out is left blank
                log.info("core2 setup check: no figure for %s: %s", row["slug"], str(exc)[:160])
        priority = "Please review" if row["notes"] else "Optional"
        put(ws, r, [priority, row["shown"], row["works_out"], row["adds_up"] or "—", SHOWN_AS.get(row["format"],
                    row["format"]), row["date"] or "None found",
                    f"{_shown_value(figure, row['format'])} {said}".strip() if figure is not None else "—",
                    " ".join(row["notes"]) or "—", None, None, None, ", ".join(row["synonyms"]) or None, None,
                    row["id"]],
            answers=((4,) if row["adds_up"] else ()) + (5, 6, 9, 10, 11, 12, 13), priority=priority)
        if row["adds_up"]:
            choices(ws, f"D{r}", list(ADD_UP.values()))
        if row["dates"]:
            choices(ws, f"F{r}", [n for n, _ in row["dates"]])
        r += 1
    for s in suggestions or []:
        put(ws, r, ["Suggested", s["name"], s.get("works_out") or "", "—", SHOWN_AS.get(s.get("format", ""), ""),
                    s.get("date") or "", s.get("figure") or "—", "Solviz suggests adding this metric. Pick Yes in "
                    "\"Right?\" to add it, No to drop it.", None, None, None, None, None, _id("suggest", s["id"])],
            answers=(9, 10), priority="Suggested")
        r += 1
    choices(ws, f"E{FIRST_ROW}:E{r}", list(SHOWN_AS.values()))
    choices(ws, f"I{FIRST_ROW}:I{r}", YES_NO)
    choices(ws, f"M{FIRST_ROW}:M{r}", ["Yes", "No"])
    for _ in range(4):
        put(ws, r, ["", "(A metric your team uses)", "(How it is worked out, in words)", "", "", "", "", "", None, None,
                    None, None, None, ""], answers=(2, 3, 10, 11))
        r += 1

    # 3 Tables and dates
    ws = sheet("3 Tables and dates", ["Priority", "Table", "One row is", "Rows", "Counted by date", "Right?",
                                      "Your correction", "Technical (for IT)", ID_HEADER],
               [15, 20, 42, 11, 22, 10, 34, 30, 10],
               "Business owner or IT. Which date each table is counted by: an order by its order date, a payment by "
               "its payment date.")
    r = FIRST_ROW
    for t in model.tables.values():
        if t.kind == "calendar" or t.hidden:
            continue
        dates = _dates(model, t.key)
        priority = "Please review" if t.kind in ("fact", "snapshot") else "Optional"
        put(ws, r, [priority, _table(model, t.key), t.grain_text or t.description or "",
                    f"{t.row_count:,}" if t.row_count else "", _date_shown(model, t.key),
                    None, None, _physical(t.key), _id("table", t.key)], answers=(3, 5, 6, 7), quiet=(8,),
            priority=priority)
        if dates:
            choices(ws, f"E{r}", [n for n, _ in dates])
        r += 1
    choices(ws, f"F{FIRST_ROW}:F{r}", YES_NO)

    # 4 Links
    ws = sheet("4 Links", ["Priority", "From", "To", "Matched on", "What it means", "Rows that find a match",
                           "Rows without a match", "Right?", "Your correction", "Technical (for IT)", ID_HEADER],
               [15, 20, 20, 26, 40, 14, 18, 10, 34, 40, 10],
               "IT team. How the tables link. A wrong link makes figures wrong, so check the rows marked Please review.")
    r = FIRST_ROW
    joins = [j for j in model.joins.values() if not j.to_calendar and j.trust != "rejected"
             and j.from_table in model.tables and j.to_table in model.tables]
    for j in sorted(joins, key=lambda j: (j.match_rate >= 0.98 and j.trust != "proposed", _table(model, j.from_table))):
        f, t = _table(model, j.from_table), _table(model, j.to_table)
        meaning = {"many_to_one": f"Each {f.lower()} belongs to one {t.lower()}",
                   "one_to_one": f"Each {f.lower()} has exactly one {t.lower()}"}.get(
            j.cardinality, f"{_plural(f)} and {_plural(t).lower()} meet many to many")
        if j.role:
            meaning += f" (as {j.role})"
        priority = "Please review" if j.match_rate < 0.98 or j.trust == "proposed" else "Optional"
        put(ws, r, [priority, f, t, ", ".join(f"{_column(model, a)} = {_column(model, b)}"
                                              for a, b in zip(j.from_columns, j.to_columns)), meaning,
                    f"{100 * j.match_rate:.1f}%", KEPT if j.keep_unmatched else LEFT_OUT, None, None,
                    j.key.replace("memory.main.", ""), _id("link", j.key)], answers=(7, 8, 9), quiet=(10,),
            priority=priority)
        r += 1
    choices(ws, f"G{FIRST_ROW}:G{r}", [KEPT, LEFT_OUT])
    choices(ws, f"H{FIRST_ROW}:H{r}", YES_NO)

    # 5 Codes
    ws = sheet("5 Codes", ["Priority", "Where", "Code", "Rows", "Solviz shows it as", "What it means", "Technical (for IT)",
                           ID_HEADER], [15, 30, 12, 11, 30, 34, 34, 10],
               "Business owner or IT. Short codes and what they mean, so answers show names, not codes.")
    r = FIRST_ROW
    for row in code_rows(model):
        priority = "Optional" if row["named"] else "Please review"
        put(ws, r, [priority, row["where"], row["code"], f"{row['rows']:,}", row["named"] or f"{row['code']} (the code)",
                    None, row["technical"], row["id"]], answers=(6,), quiet=(7,), priority=priority)
        r += 1

    # 6 Your words: "It means" picks a metric or field, or says it in words
    lists = wb.create_sheet(LISTS)
    names = sorted({m.business_name for m in model.measures.values() if not m.hidden}
                   | {a.business_name for a in model.attributes.values()}
                   | {e.business_name for e in model.entities.values()}, key=str.casefold)
    for i, name in enumerate(names, start=1):
        lists.cell(row=i, column=1, value=name).data_type = "s"
    lists.sheet_state = "hidden"
    ws = sheet("6 Your words", ["Word or phrase your team uses", "It means", "Example question", "Note"],
               [30, 40, 50, 30], "Business owner. Words that mean something specific in your business. Pick what "
               "it means from the list, or write it in your own words.")
    put(ws, FIRST_ROW, ["EXAMPLE: book", "(a metric or field from the list)", "What did we book in Q2 by branch?",
                        "Example: replace or delete"], quiet=(1, 2, 3, 4))
    for r in range(FIRST_ROW + 1, FIRST_ROW + 26):
        put(ws, r, [None, None, None, None], answers=(1, 2, 3, 4))
    if names:
        choices(ws, f"B{FIRST_ROW + 1}:B{FIRST_ROW + 25}", f"{LISTS}!$A$1:$A${len(names)}", strict=False)

    # 7 Test questions
    ws = sheet("7 Test questions", ["Question your team asks", "The answer in your own report", "Which report",
                                    "Period", "Note"], [50, 26, 24, 16, 30],
               "Business owner. 20 to 50 real questions with the answer from your own reports. Solviz is checked "
               "against them before your team uses it, and again after every change.")
    put(ws, FIRST_ROW, ["EXAMPLE: How many new orders did each region take in Q2?", "North: 41",
                        "Monthly sales pack", "Q2", "Example: replace or delete"], quiet=(1, 2, 3, 4, 5))
    for r in range(FIRST_ROW + 1, FIRST_ROW + 51):
        put(ws, r, [None, None, None, None, None], answers=(1, 2, 3, 4, 5))

    meta = wb.create_sheet(META)
    for i, (k, v) in enumerate([("version", VERSION), ("workspace", account_id), ("exported", today.isoformat()),
                                ("month", month)], start=1):
        meta.cell(row=i, column=1, value=k)
        meta.cell(row=i, column=2, value=v)
    meta.sheet_state = "hidden"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── reading the answers ─────────────────────────────────────────────────────

def read(data: bytes, *, account_id: str) -> dict[str, list[dict[str, Any]]]:
    """The answered rows of each tab, as {header: value}, from a setup check exported for this workspace."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - anything openpyxl cannot open is not a setup check
        raise SetupCheckError("This file cannot be read as an Excel workbook (.xlsx).") from None
    if META not in wb.sheetnames:
        raise SetupCheckError("This is not a Solviz setup check: download one from this page and fill that in.")
    meta = {str(r[0]): r[1] for r in wb[META].iter_rows(values_only=True) if r and r[0]}
    if meta.get("version") != VERSION:
        raise SetupCheckError("This setup check is from another version of Solviz: download a new one.")
    if str(meta.get("workspace")) != account_id:
        raise SetupCheckError("This setup check was made for another workspace.")
    out: dict[str, list[dict[str, Any]]] = {"_month": [{"month": meta.get("month") or ""}]}
    for ws in wb.worksheets:
        if ws.title in (META, LISTS, "Start here"):
            continue
        rows = ws.iter_rows(min_row=HEAD_ROW, values_only=True)
        headers = [str(h or "").strip() for h in next(rows, ())]
        answered = []
        for values in rows:
            row = {h: (v.strip() if isinstance(v, str) else v) for h, v in zip(headers, values) if h}
            if any(v not in (None, "") for v in row.values()):
                answered.append(row)
        out[ws.title] = answered
    return out


# ── what the answers would change ───────────────────────────────────────────

def _item(tab: str, what: str, change: str, ops: list[dict[str, Any]], check: str = "", ok: bool = True) -> dict:
    if not check and any(o["op"] == "request" for o in ops):
        check = next(o["note"] for o in ops if o["op"] == "request")     # the team's own words, for the admin
    return {"tab": tab, "what": what, "change": change, "ops": ops, "check": check, "ok": ok}


def _request(target_kind: str, target_key: str, target_name: str, note: str) -> dict[str, Any]:
    return {"op": "request", "target_kind": target_kind, "target_key": target_key, "target_name": target_name,
            "note": note}


def _new_metric(name: str, words: str, example: str = "") -> dict[str, Any]:
    """A metric the team uses and Solviz has not: a New metric request in their own name and words, which the
    Requests page opens in the metric editor, written from those words."""
    return {"op": "request", "kind": "new_metric", "target_kind": "measure", "target_key": "",
            "target_name": name[:120], "note": " ".join(words.split())[:1000], "example": example[:300]}


def _ident(row: dict[str, Any]) -> list[Any] | None:
    try:
        parts = json.loads(str(row.get(ID_HEADER) or ""))
    except ValueError:
        return None
    return parts if isinstance(parts, list) and parts else None


def _blank(value: Any) -> bool:
    return value is None or str(value).strip() in ("", "—")


def review(model: SemanticModel, answers: dict[str, list[dict[str, Any]]], *, warehouse: Any = None,
           today: dt.date | None = None) -> list[dict[str, Any]]:
    """The changes the client's answers propose, in words, each checked on the data. Nothing is written."""
    from core2.model.links import check_table, value_problem
    from core2.model.overrides import target

    today = today or dt.date.today()
    items: list[dict[str, Any]] = []

    # 1 Rows to leave out
    for row in answers.get("1 Rows to leave out", []):
        ident, choice = _ident(row), str(row.get("Leave them out of totals?") or "")
        note = str(row.get("Your note") or "").strip()
        if ident is None:
            text = str(row.get("These rows") or "").strip()
            if text and not text.startswith("(") or note:
                items.append(_item("Rows to leave out", text or "Other rows", "To the Requests queue",
                                   [_request("table", "", "", f"Setup check, rows to leave out: {text}. {note}".strip())]))
            continue
        kind, table = ident[0], ident[1]
        if table not in model.tables:
            items.append(_item("Rows to leave out", str(row.get("These rows")), "Not applied", [],
                               "This table is not in the data any more.", False))
            continue
        if kind == "leave" and choice == "Yes":
            column, value = ident[2], ident[3]
            if column not in model.columns:
                items.append(_item("Rows to leave out", str(row.get("These rows")), "Not applied", [],
                                   "This field is not in the data any more.", False))
                continue
            rule = ColumnFilter(column=column, op="ne", values=[_typed(model, column, value)])
            problem = value_problem(model, rule)
            check, ok = (problem, False) if problem else ("", True)
            if ok and warehouse is not None:
                found = check_table(model, table, [rule], warehouse, today)
                if found.get("problem"):
                    check, ok = found["problem"], False
                else:
                    check = (f"Leaves out {found['left_out']:,} of {found['rows']:,} "
                             f"{_plural(_table(model, table)).lower()}")
                    if found.get("measure") and found.get("amount") is not None:
                        check += f" ({_shown_value(found['amount'], found.get('format', 'number'))} of " \
                                 f"{found['measure']}{' in ' + found['year'] if found.get('year') else ''})"
            items.append(_item("Rows to leave out", str(row.get("These rows")), "Leave them out of every total",
                               [{"op": "add_filter", "table": table, "filter": rule.model_dump()}], check, ok))
        elif kind == "leftout" and choice == "No":
            rule = next((f for f in model.tables[table].default_filters if _filter_id(f) == ident[2]), None)
            if rule is None:
                continue                                    # already taken off since the workbook was made
            items.append(_item("Rows to leave out", str(row.get("These rows")), "Count them in totals again",
                               [{"op": "remove_filter", "table": table, "filter": rule.model_dump()}]))
        if note and choice != "Yes" or choice == "Not sure":
            items.append(_item("Rows to leave out", str(row.get("These rows")), "To the Requests queue",
                               [_request("table", table, _table(model, table),
                                         f"Setup check, {row.get('These rows')}: {choice or 'no answer'}. {note}")]))

    # 2 Metrics
    month = (answers.get("_month") or [{}])[0].get("month") or ""
    current = {r["id"]: r for r in metric_rows(model)}
    for row in answers.get("2 Metrics", []):
        ident = _ident(row)
        right = str(row.get("Right?") or "")
        should = str(row.get("What it should be") or "").strip()
        if ident is None:
            name, how = str(row.get("Metric") or "").strip(), str(row.get("What Solviz works out") or "").strip()
            how = "" if how.startswith("(") else how
            if name and not name.startswith("(") and (how or should):
                items.append(_item("Metrics", name, "To the Requests queue, as a new metric to add",
                                   [_new_metric(name, f"{how} {should}")]))
            continue
        if ident[0] == "suggest":
            if right in ("Yes", "No"):
                items.append(_item("Metrics", str(row.get("Metric")), "Add this metric" if right == "Yes"
                                   else "Do not add this metric",
                                   [{"op": "accept_request" if right == "Yes" else "reject_request",
                                     "id": int(ident[1])}]))
            continue
        key = ident[1]
        measure, now = model.measures.get(key), current.get(_id("metric", key))
        if measure is None or now is None:
            items.append(_item("Metrics", str(row.get("Metric")), "Not applied", [],
                               "This metric is not in the data any more.", False))
            continue
        obj = target("measure", key)
        adds = str(row.get("Add up as") or "")
        if now["adds_up"] and adds in ("Total", "Average") and adds != now["adds_up"]:
            items.append(_item("Metrics", measure.business_name, f"Add up as: {adds} (was {now['adds_up']})",
                               [{"op": "set", "object_key": obj, "field": "additivity",
                                 "value": "non_additive" if adds == "Average" else "additive"}]))
        shown = str(row.get("Shown as") or "")
        fmt = next((k for k, v in SHOWN_AS.items() if v == shown), None)
        if fmt and fmt != measure.format:
            items.append(_item("Metrics", measure.business_name, f"Shown as: {shown} (was "
                               f"{SHOWN_AS.get(measure.format, measure.format)})",
                               [{"op": "set", "object_key": obj, "field": "format", "value": fmt}]))
        date = str(row.get("Counted by date") or "")
        if date and date != now["date"] and date != "None found":
            role = next((k for n, k in now["dates"] if n == date), None)
            items.append(_item("Metrics", measure.business_name, f"Counted by date: {date} (was {now['date'] or 'none'})",
                               [{"op": "set", "object_key": obj, "field": "default_date", "value": role}] if role else [],
                               "" if role else f"{date} is not a date of {_table(model, measure.table)}.", bool(role)))
        words = [w for w in _words(row.get("Other names your team uses")) if w not in now["synonyms"]]
        if words:
            items.append(_item("Metrics", measure.business_name, "Other names: " + ", ".join(words),
                               [{"op": "add_synonyms", "object_key": obj, "words": words}]))
        if str(row.get("Hide it?") or "") == "Yes":
            items.append(_item("Metrics", measure.business_name, "Hide it from readers",
                               [{"op": "set", "object_key": obj, "field": "hidden", "value": True}]))
        own_text = row.get(next((h for h in row if h.startswith("Your own figure")), ""))
        shown = row.get(next((h for h in row if h.startswith("Solviz's figure")), "")) or "—"
        own, theirs = _number(own_text), _number(shown)
        differs = False
        if own is not None:
            if theirs is not None and abs(own - theirs) <= max(0.005 * abs(own), 0.5):
                items.append(_item("Metrics", measure.business_name, "Nothing to change", [],
                                   f"The team's figure for {month} ({own_text}) matches Solviz's ({shown})."))
            else:
                differs = True
                items.append(_item("Metrics", measure.business_name, "To the Requests queue",
                                   [_request("measure", key, measure.business_name,
                                             f"Setup check: the team's own figure for {month} is {own_text}; Solviz "
                                             f"shows {shown}. {should}".strip())],
                                   f"The team's figure for {month} ({own_text}) is not Solviz's ({shown}): find out "
                                   "why before the team uses it.", False))
        if right in ("No", "Not sure") or should and not differs:
            items.append(_item("Metrics", measure.business_name, "To the Requests queue",
                               [_request("measure", key, measure.business_name,
                                         f"Setup check: {right or 'a correction'}. {should}".strip())]))

    # 3 Tables and dates
    for row in answers.get("3 Tables and dates", []):
        ident = _ident(row)
        if ident is None or ident[1] not in model.tables:
            continue
        table = model.tables[ident[1]]
        dates = _dates(model, table.key)
        date = str(row.get("Counted by date") or "")
        now_date = _date_shown(model, table.key)
        if date and date not in (now_date, "None found"):
            role = next((k for n, k in dates if n == date), None)
            items.append(_item("Tables and dates", _table(model, table.key), f"Counted by date: {date} (was {now_date})",
                               [{"op": "default_date", "table": table.key, "role": role}] if role else [],
                               "" if role else f"{date} is not a date of {_table(model, table.key)}.", bool(role)))
        grain = str(row.get("One row is") or "").strip()
        if grain and grain != (table.grain_text or table.description or "").strip():
            items.append(_item("Tables and dates", _table(model, table.key), f"One row is: {grain}",
                               [{"op": "set", "object_key": target("table", table.key), "field": "grain_text",
                                 "value": grain[:200]}]))
        right, fix = str(row.get("Right?") or ""), str(row.get("Your correction") or "").strip()
        if right in ("No", "Not sure") or fix:
            items.append(_item("Tables and dates", _table(model, table.key), "To the Requests queue",
                               [_request("table", table.key, _table(model, table.key),
                                         f"Setup check: {right or 'a correction'}. {fix}".strip())]))

    # 4 Links
    for row in answers.get("4 Links", []):
        ident = _ident(row)
        if ident is None:
            text = " ".join(str(row.get(h) or "") for h in ("From", "To", "Matched on", "Your correction")).strip()
            if text and not text.startswith("("):
                items.append(_item("Links", text[:80], "To the Requests queue",
                                   [_request("join", "", "", f"Setup check, a link to add: {text}")]))
            continue
        join = model.joins.get(ident[1])
        if join is None:
            continue
        name = f"{_table(model, join.from_table)} → {_table(model, join.to_table)}"
        keep = str(row.get("Rows without a match") or "")
        if keep in (KEPT, LEFT_OUT) and (keep == KEPT) != join.keep_unmatched:
            obj = target("join", join.key)
            items.append(_item("Links", name, f"Rows without a match: {keep}",
                               [{"op": "set", "object_key": obj, "field": "keep_unmatched", "value": keep == KEPT},
                                {"op": "set", "object_key": obj, "field": "trust", "value": "admin"}]))
        right, fix = str(row.get("Right?") or ""), str(row.get("Your correction") or "").strip()
        if right in ("No", "Not sure") or fix:
            items.append(_item("Links", name, "To the Requests queue",
                               [_request("join", join.key, name, f"Setup check: {right or 'a correction'}. {fix}".strip())]))

    # 5 Codes
    for row in answers.get("5 Codes", []):
        ident, meaning = _ident(row), str(row.get("What it means") or "").strip()
        if ident is None or not meaning or ident[1] not in model.columns:
            continue
        column, code = ident[1], str(ident[2])
        items.append(_item("Codes", f"{row.get('Where')}: {code}", f"Show {code} as {meaning}",
                           [{"op": "name_value", "column": column, "code": code, "name": meaning[:60]}]))

    # 6 Your words
    by_name: dict[str, tuple[str, str]] = {}
    for m in model.measures.values():
        if not m.hidden:
            by_name.setdefault(m.business_name.casefold(), ("measure", target("measure", m.key)))
    for a in model.attributes.values():
        if a.column in model.columns:
            by_name.setdefault(a.business_name.casefold(), ("field", target("column", a.column)))
    for e in model.entities.values():
        by_name.setdefault(e.business_name.casefold(), ("thing", target("entity", e.slug)))
    for row in answers.get("6 Your words", []):
        word, means = str(row.get("Word or phrase your team uses") or "").strip(), str(row.get("It means") or "").strip()
        if not word or word.upper().startswith("EXAMPLE"):
            continue
        found = by_name.get(means.casefold())
        if found:
            items.append(_item("Your words", word, f"\"{word}\" also means {means}",
                               [{"op": "add_synonyms", "object_key": found[1], "words": [word[:40]]}]))
        elif means:
            # A word for something Solviz has no name for ("GDR: generic scripts as a % of all scripts"): a metric
            # the team uses, to be written in the editor from their words.
            items.append(_item("Your words", word, "To the Requests queue, as a new metric to add",
                               [_new_metric(word, means, str(row.get("Example question") or ""))]))
        else:
            items.append(_item("Your words", word, "Not applied", [], "No meaning given.", False))

    # 7 Test questions
    for row in answers.get("7 Test questions", []):
        question = str(row.get("Question your team asks") or "").strip()
        if not question or question.upper().startswith("EXAMPLE"):
            continue
        expected = str(row.get("The answer in your own report") or "").strip()
        items.append(_item("Test questions", question[:120], "Kept as a test question",
                           [{"op": "test_question", "question": question[:500], "expected": expected[:300],
                             "report": str(row.get("Which report") or "")[:120],
                             "period": str(row.get("Period") or "")[:60], "note": str(row.get("Note") or "")[:300]}],
                           "" if expected else "No answer given: it can be run, not checked."))

    for i, item in enumerate(items):
        item["n"] = i
    return items


# ── writing the chosen changes ──────────────────────────────────────────────

def apply(account_id: str, client: dict[str, Any], model: SemanticModel, items: list[dict[str, Any]], *, by: str,
          check_id: int, file_name: str = "") -> dict[str, int]:
    """Write the chosen items: decisions as the Data pages write them, requests to the queue, test questions kept."""
    import store
    from core2 import request_service
    from core2.model.overrides import target

    db_id = client.get("db_config_id")
    note = f"Setup check #{check_id}{' (' + file_name + ')' if file_name else ''}"
    counts = {"decisions": 0, "requests": 0, "test_questions": 0, "metrics_added": 0, "skipped": 0}
    filters = {k: [f.model_dump() for f in t.default_filters] for k, t in model.tables.items()}
    touched_filters: set[str] = set()
    synonyms: dict[str, list[str]] = {}
    names: dict[str, dict[str, str]] = {}

    def write(object_key: str, field: str, value: Any) -> None:
        store.set_core2_override(account_id, db_id, object_key, field, value, author=by, note=note)
        counts["decisions"] += 1

    for item in items:
        for op in item.get("ops") or []:
            kind = op.get("op")
            if kind == "add_filter":
                if op["filter"] not in filters.setdefault(op["table"], []):
                    filters[op["table"]].append(op["filter"])
                touched_filters.add(op["table"])
            elif kind == "remove_filter":
                filters[op["table"]] = [f for f in filters.get(op["table"], []) if f != op["filter"]]
                touched_filters.add(op["table"])
            elif kind == "set":
                write(op["object_key"], op["field"], op["value"])
            elif kind == "default_date":
                for role in model.date_roles.values():
                    if role.table == op["table"]:
                        store.delete_core2_override(account_id, db_id, target("date_role", role.key), "is_default")
                write(target("table", op["table"]), "default_date", op["role"])
            elif kind == "add_synonyms":
                synonyms.setdefault(op["object_key"], []).extend(op["words"])
            elif kind == "name_value":
                names.setdefault(op["column"], {})[op["code"]] = op["name"]
            elif kind == "request":
                store.add_core2_request(account_id, db_id, kind=op.get("kind", "change"), user_id="",
                                        user_name=f"Setup check ({by})", target_kind=op.get("target_kind", ""),
                                        target_key=op.get("target_key", ""), target_name=op.get("target_name", ""),
                                        note=op["note"][:1000], example=op.get("example", ""), source="setup_check")
                counts["requests"] += 1
            elif kind in ("accept_request", "reject_request"):
                request = store.get_core2_request(account_id, int(op["id"]))
                if not request or request.get("status") != "waiting":
                    counts["skipped"] += 1
                    continue
                try:
                    if kind == "accept_request":
                        request_service.accept(account_id, client, request, decided_by=by)
                        counts["metrics_added"] += 1
                    else:
                        request_service.reject(account_id, request, decided_by=by,
                                               reason="Not wanted, said the setup check")
                except Exception as exc:  # noqa: BLE001 - one refused request never stops the rest
                    log.warning("core2 setup check: request %s not decided: %s", op["id"], exc)
                    counts["skipped"] += 1
            elif kind == "test_question":
                store.add_core2_test_question(account_id, db_id, question=op["question"], expected=op["expected"],
                                              report=op["report"], period=op["period"], note=op["note"],
                                              setup_check_id=check_id)
                counts["test_questions"] += 1

    for table in touched_filters:
        write(target("table", table), "default_filters", filters[table])
    for object_key, words in synonyms.items():
        kind, _, key = object_key.partition(":")
        thing: Any = {"measure": model.measures, "column": model.columns, "entity": model.entities}[kind].get(key)
        have = dict(thing.synonyms) if thing is not None else {}
        have["en"] = list(dict.fromkeys([*have.get("en", []), *words]))[:20]
        write(object_key, "synonyms", have)
    for column, named in names.items():
        have = dict(getattr(model.columns.get(column), "value_names", {}) or {})
        write(target("column", column), "value_names", {**have, **named})
    return counts
