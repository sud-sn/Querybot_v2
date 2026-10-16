"""A workbook is a workspace's database: chosen on the Databases page, discovered, learned and asked.

An admin picks "Excel workbook", uploads a file (or pastes a SharePoint or OneDrive link) and saves it as
any other connection. The connection test names the file and its sheets; discovery writes each sheet as a
table, its headers as the columns' descriptions; Learn learns it as it learns any warehouse; questions are
compiled in DuckDB's SQL and read the sheets' own figures. Nothing a question sends can change the file.

The synthetic purchasing warehouse, written out as a workbook (one sheet per table, a title above one
sheet's header), and its golden questions with their reference SQL.
"""

from __future__ import annotations

import asyncio
import datetime as dt
import io
import json
from unittest.mock import patch

import httpx
import pytest
from openpyxl import Workbook
from starlette.requests import Request

from core import workbook_source as W
from tests.test_snowflake_signs_in_as_a_service_user import _request, admin  # noqa: F401 - the fixture


def as_workbook(built) -> bytes:
    """Each table of a built warehouse as a sheet, headers as people write them ("Order Date")."""
    book = Workbook(write_only=True)
    for n, (table, physical) in enumerate(sorted(built.tables.items())):
        sheet = book.create_sheet(table.replace("_", " ").title()[:31])
        cursor = built.con.execute(f'SELECT * FROM "{physical}"')
        names = [d[0] for d in cursor.description]
        if n == 0:
            sheet.append([f"{table.replace('_', ' ').title()} export"])
            sheet.append([])
        sheet.append([name.replace("_", " ").title() for name in names])
        for row in cursor.fetchall():
            sheet.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture(scope="module")
def purchasing(tmp_path_factory):
    from evals.core2 import domains
    from evals.core2.framework import materialize

    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("QUERYBOT_WORKBOOKS_DIR", str(tmp_path_factory.mktemp("workbooks")))
        built = materialize(domains.build("purchasing"), "descriptive")
        loaded = W.save(as_workbook(built), "Purchasing.xlsx")
        yield built, {"source": W.UPLOAD, "workbook": loaded.sha256, "file_name": "Purchasing.xlsx"}


@pytest.fixture(scope="module")
def learned(purchasing, tmp_path_factory):
    """Learn on the workbook, from what discovery wrote, as production runs it."""
    from core.schema import discover_and_write
    from core2.bootstrap.build import BuildOptions, build_model
    from core2.bootstrap.inventory import from_schema_json
    from core2.warehouse.querybot import QueryBotWarehouse

    built, creds = purchasing
    folder = tmp_path_factory.mktemp("xl_schema")
    count = discover_and_write(creds, "excel", str(folder))
    schema = json.loads((folder / "_schema.json").read_text())
    inventory = from_schema_json(schema, "excel")
    lines: list[str] = []
    with QueryBotWarehouse("excel", creds) as warehouse:
        model = build_model(warehouse, inventory, client_id="xl-test",
                            options=BuildOptions(workers=2, progress=lines.append))
    return built, creds, model, lines, count, schema, folder


def test_the_connection_test_names_the_file_and_its_sheets(purchasing):
    from core.schema import test_connection

    built, creds = purchasing
    found = test_connection(creds, "excel")
    assert found["workbook"] == "Purchasing.xlsx"
    for table in built.tables:
        assert f"{table} (" in found["tables"]
    assert "write_permissions" not in found


def test_discovery_writes_each_sheet_as_a_table_with_its_headers(learned):
    built, _creds, _model, _lines, count, schema, folder = learned
    tables = {k: v for k, v in schema.items() if not k.startswith("__")}
    assert count == len(built.tables) == len(tables)
    lines = tables["WORKBOOK.MAIN.purchase_order_lines"]
    assert (lines["database"], lines["schema"]) == ("workbook", "main")
    assert lines["row_count"] == built.con.execute("SELECT COUNT(*) FROM purchase_order_lines").fetchone()[0]
    header = {c["name"]: c["comment"] for c in lines["columns"]}
    assert header["order_date"] == "Order Date" and header["line_amount"] == "Line Amount"
    assert "Purchase Order Lines" in lines["comment"]
    assert (folder / "purchase_order_lines.md").exists()


def test_learn_on_a_workbook_learns_its_tables_and_links(learned):
    built, _creds, model, lines, *_ = learned
    assert {t.name for t in model.tables.values()} == set(built.tables)
    links = {(model.tables[j.from_table].name, model.tables[j.to_table].name) for j in model.joins.values()}
    assert ("purchase_order_lines", "suppliers") in links and ("purchase_order_lines", "buyers") in links
    assert not [line for line in lines if line.startswith("Left out ")], lines


def test_every_golden_question_gives_the_reference_rows_from_a_workbook(learned):
    from core2.compile.compiler import compile_query
    from core2.resolve.resolver import Context, resolve
    from core2.warehouse.querybot import QueryBotWarehouse
    from core2.warehouse.runner import DuckDBWarehouse
    from evals.core2.compile_eval import Untranslatable, _Words, golden, same_rows

    built, creds, model, *_ = learned
    cases = golden(built.domain.name)
    today = dt.date.fromisoformat(str(cases["today"]))
    words = _Words(model, built)
    reference = DuckDBWarehouse(built.con)
    ran, wrong = 0, []
    with QueryBotWarehouse("excel", creds) as warehouse:
        assert warehouse.dialect == "duckdb"
        for case in cases["questions"]:
            if not case.get("plan") or not case.get("reference_sql"):
                continue
            try:
                logical = resolve(words.plan(case["plan"]), model, Context(today=today))
            except Untranslatable:
                continue
            sql = compile_query(logical, model, warehouse.dialect).sql
            got = warehouse.query(sql)
            ran += 1
            expected = reference.query(case["reference_sql"])
            diff = same_rows(expected.columns, expected.rows, got.columns, got.rows,
                             order_matters=bool((case.get("expect") or {}).get("order_matters")))
            if diff:
                wrong.append(f"{case['id']}: {diff}\n{sql}")
    assert ran >= 5
    assert not wrong, "\n\n".join(wrong)


@pytest.mark.parametrize("sql", ["DELETE FROM workbook.main.buyers",
                                 "CREATE TABLE workbook.main.x AS SELECT 1 AS a",
                                 "SELECT * FROM read_csv('/etc/hosts')",
                                 "COPY workbook.main.buyers TO '/tmp/leak.csv'"])
def test_nothing_a_query_sends_changes_the_workbook_or_reads_another_file(purchasing, sql):
    from core.schema import run_query
    from core2.warehouse.querybot import QueryBotWarehouse

    _built, creds = purchasing
    with pytest.raises(Exception):
        run_query(creds, "excel", sql)
    with QueryBotWarehouse("excel", creds) as warehouse, pytest.raises(Exception):
        warehouse.query(sql)
    assert run_query(creds, "excel", "SELECT COUNT(*) AS n FROM workbook.main.buyers")[0]["n"] == 5


# ── The Databases page ───────────────────────────────────────────────────────


def _multipart(path: str, files: dict, data: dict | None = None) -> Request:
    built = httpx.Request("POST", f"http://testserver{path}", files=files, data=data or {})
    body = built.read()

    async def receive():
        return {"type": "http.request", "body": body, "more_body": False}

    return Request({
        "type": "http", "method": "POST", "path": path, "root_path": "", "scheme": "http", "query_string": b"",
        "server": ("testserver", 80), "client": ("127.0.0.1", 1),
        "headers": [(b"content-type", built.headers["content-type"].encode()),
                    (b"content-length", str(len(body)).encode())]}, receive)


def _upload(routes, data: bytes, name: str = "Sales.xlsx") -> dict:
    with patch.object(routes, "_is_auth", return_value=True):
        response = asyncio.run(routes.database_workbook_upload(_multipart(
            "/admin/databases/workbook/upload", {"file": (name, data, "application/octet-stream")})))
    return json.loads(response.body)


def _save(routes, fields: dict):
    with patch.object(routes, "_is_auth", return_value=True):
        return asyncio.run(routes.database_save(_request("/admin/databases/save", fields), routes.BackgroundTasks()))


def _page(routes) -> str:
    import html

    with patch.object(routes, "_is_auth", return_value=True):
        page = asyncio.run(routes.databases_page(_request("/admin/databases", method="GET"))).body.decode()
    return html.unescape(page)


def _small_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.append(["Order ID", "Order Date", "Amount"])
    sheet.append([1, dt.date(2025, 1, 5), 12.5])
    sheet.append([2, dt.date(2025, 2, 9), 30.0])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture()
def workbooks(tmp_path, monkeypatch):
    monkeypatch.setenv("QUERYBOT_WORKBOOKS_DIR", str(tmp_path / "workbooks"))
    W._LINKS.clear()


def test_the_form_offers_a_workbook_and_reads_its_fields(admin, workbooks):  # noqa: F811
    from admin.routes import _db_credentials_from_form

    page = _page(admin)
    assert '<option value="excel">Excel workbook</option>' in page
    assert 'name="xl_link"' in page and 'id="xl_file"' in page
    _, name, _, upload = _db_credentials_from_form({"db_type": "excel", "name": "Sales", "xl_source": "upload",
                                                    "xl_workbook": "ab" * 32, "xl_file_name": "Sales.xlsx",
                                                    "xl_link": "ignored"})
    assert upload["source"] == "upload" and upload["workbook"] == "ab" * 32 and "link" not in upload
    _, _, _, link = _db_credentials_from_form({"db_type": "excel", "name": "Sales", "xl_source": "link",
                                               "xl_link": " https://contoso.sharepoint.com/:x:/s/f/E1 ",
                                               "xl_graph": "app", "xl_tenant_id": "t", "xl_client_id": "c",
                                               "xl_client_secret": "s3cret"})
    assert {k: link[k] for k in ("source", "link", "graph", "tenant_id", "client_id", "client_secret")} == {
        "source": "link", "link": "https://contoso.sharepoint.com/:x:/s/f/E1", "graph": "app",
        "tenant_id": "t", "client_id": "c", "client_secret": "s3cret"}
    assert "log_export_enabled" not in link or link["log_export_enabled"] == "0"


def test_an_uploaded_workbook_is_read_saved_listed_and_kept(admin, workbooks):  # noqa: F811
    told = _upload(admin, _small_workbook())
    assert told["status"] == "ok" and told["file_name"] == "Sales.xlsx"
    assert [(t["name"], t["rows"]) for t in told["tables"]] == [("orders", 2)]
    response = _save(admin, {"db_type": "excel", "name": "Sales workbook", "xl_source": "upload",
                             "xl_workbook": told["workbook"], "xl_file_name": told["file_name"]})
    assert response.status_code == 303 and "error" not in response.headers["location"]
    saved = admin.store.list_db_configs()
    assert [(d["db_type"], d["label"], d["dialect"]) for d in saved] == [("excel", "Excel workbook", "duckdb")]
    page = _page(admin)
    assert "Sales.xlsx" in page
    # Saved again without a new file: the one stored is kept.
    _save(admin, {"db_type": "excel", "name": "Sales workbook", "db_id": str(saved[0]["id"]), "xl_source": "upload"})
    assert admin.store.get_db_config(saved[0]["id"])["credentials"]["workbook"] == told["workbook"]


def test_a_file_that_is_no_workbook_is_refused_with_the_reason(admin, workbooks):  # noqa: F811
    told = _upload(admin, b"order_id,qty\n1,2\n", "sales.csv")
    assert told["status"] == "error" and "not an Excel workbook" in told["message"]
    response = _save(admin, {"db_type": "excel", "name": "Sales", "xl_source": "upload", "xl_workbook": "00" * 32})
    assert "error=" in response.headers["location"]          # a workbook never uploaded is not saved


def test_a_linked_workbooks_app_secret_is_never_shown_and_is_kept(admin, workbooks):  # noqa: F811
    fields = {"db_type": "excel", "name": "Shared sales", "xl_source": "link",
              "xl_link": "https://contoso.sharepoint.com/:x:/s/finance/EaBcD123", "xl_graph": "app",
              "xl_tenant_id": "tenant-1", "xl_client_id": "app-1", "xl_client_secret": "s3cret-value"}
    response = _save(admin, fields)
    assert response.status_code == 303 and "error" not in response.headers["location"]
    saved = admin.store.list_db_configs()[0]
    page = _page(admin)
    assert "s3cret-value" not in page and "contoso.sharepoint.com" in page
    row = page.split('data-db-type="excel"', 1)[1].split(">", 1)[0]
    edit = json.loads(row.split("data-db-creds='", 1)[1].split("'", 1)[0])
    assert edit["link"] == fields["xl_link"] and edit["tenant_id"] == "tenant-1" and "client_secret" not in edit
    _save(admin, {**fields, "db_id": str(saved["id"]), "xl_client_secret": ""})
    assert admin.store.get_db_config(saved["id"])["credentials"]["client_secret"] == "s3cret-value"


def test_a_link_read_with_the_email_app_needs_no_app_of_its_own(admin, workbooks):  # noqa: F811
    response = _save(admin, {"db_type": "excel", "name": "Shared", "xl_source": "link", "xl_graph": "mail",
                             "xl_link": "https://contoso-my.sharepoint.com/:x:/p/ana/E9"})
    assert "error" not in response.headers["location"]
    missing = _save(admin, {"db_type": "excel", "name": "Shared", "xl_source": "link", "xl_graph": "app",
                            "xl_link": "https://contoso-my.sharepoint.com/:x:/p/ana/E9"})
    assert "error=" in missing.headers["location"]
