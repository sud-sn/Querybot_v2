"""The client's own team checks what Solviz learned in a workbook, and their answers become decisions.

Learn reads a database well enough to answer from it, not the business behind it: that cancelled lines never
count, what a code stands for, which rows a link should drop. The people who know will not learn the admin
console, so the model goes out as a workbook (What Solviz learned → Setup check) and comes back as changes, each
checked on the data and shown for an admin to apply. What an admin applies is an admin decision like any the Data
pages write: it changes the next answer and holds through the next Learn. Corrections only words can hold go to
the Requests queue; the team's own questions and answers are kept to check answers against.

The real routes, store and model on the invented retail warehouse; the workbook is filled in as a client would,
in its yellow cells.
"""

from __future__ import annotations

import datetime as dt
import io
import os
from unittest.mock import patch

import pytest

from evals.core2 import domains
from evals.core2.compile_eval import learn

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def retail():
    return learn(domains.build("retail"), "descriptive")


@pytest.fixture
def page(retail, tmp_path, monkeypatch):
    built, model = retail
    path = str(tmp_path / "querybot.db")
    monkeypatch.setenv("QUERYBOT_DB_PATH", path)
    monkeypatch.setenv("DB_PATH", path)
    import store
    from fastapi import FastAPI
    from starlette.testclient import TestClient

    from admin import core2_routes, core2_setup_check, routes
    from core2.warehouse.runner import DuckDBWarehouse

    store.init_db()
    account = f"acct{os.urandom(4).hex()}"
    store.upsert_client(account, "Retail (invented)")
    store.save_core2_model(account, None, model.model_dump_json())
    app = FastAPI()
    app.include_router(routes.router)
    with patch.object(routes, "_is_auth", return_value=True), \
            patch.object(core2_routes, "_is_auth", return_value=True), \
            patch.object(core2_setup_check, "_is_auth", return_value=True), \
            patch.object(core2_setup_check, "_warehouse", lambda *a: DuckDBWarehouse(built.con)):
        yield TestClient(app), account, store, DuckDBWarehouse(built.con)


def _model(account):
    from core2.bootstrap.service import load_model

    return load_model(account, None)


def _download(client, account):
    from openpyxl import load_workbook

    got = client.get(f"/admin/clients/{account}/learned/setup-check.xlsx")
    assert got.status_code == 200 and got.headers["content-type"].startswith(XLSX), got.text[:300]
    return load_workbook(io.BytesIO(got.content))


def _rows(ws):
    """Each answerable row of a tab as (row number, {header: cell})."""
    headers = [c.value for c in ws[4]]
    for r in range(5, ws.max_row + 1):
        yield r, {h: ws.cell(row=r, column=i + 1) for i, h in enumerate(headers) if h}


def _row(ws, column: str, starts: str):
    return next(cells for _, cells in _rows(ws) if str(cells[column].value or "").startswith(starts))


def _upload(client, account, wb, name="answered.xlsx"):
    buffer = io.BytesIO()
    wb.save(buffer)
    return client.post(f"/admin/clients/{account}/learned/setup-check", follow_redirects=False,
                       files={"file": (name, buffer.getvalue(), XLSX)})


def _items(store, account, check_id):
    return store.get_core2_setup_check(account, check_id)["items"]


def _figure(model, warehouse, name):
    from core2 import setup_check

    slug = next(m.slug for m in model.measures.values() if m.business_name == name)
    today = dt.date.today()
    return setup_check._figure(model, slug, warehouse, *setup_check.figures_month(model, today), today)[0]


def _month(wb) -> str:
    return {r[0]: r[1] for r in wb["_solviz"].iter_rows(values_only=True)}["month"]


def _answer(wb):
    """The team's answers: cancelled lines never count, C means Cancelled, returns without a known customer are
    dropped, net amount is also takings, turnover is gross amount, and a question with its answer."""
    _row(wb["1 Rows to leave out"], "These rows", 'Order lines with status code "C"')[
        "Leave them out of totals?"].value = "Yes"
    net = _row(wb["2 Metrics"], "Metric", "Net amount")
    net["Other names your team uses"].value = f"{net['Other names your team uses'].value or ''}, takings"
    wrong = _row(wb["2 Metrics"], "Metric", "Cost amount")
    wrong["Right?"].value, wrong["What it should be"].value = "No", "Cost should include freight."
    _row(wb["4 Links"], "From", "Return")["Rows without a match"].value = "Left out"
    code = next(c for _, c in _rows(wb["5 Codes"]) if c["Where"].value == "Order line: Status code"
                and c["Code"].value == "C")
    code["What it means"].value = "Cancelled"
    words = wb["6 Your words"]
    words["A6"], words["B6"] = "turnover", "Gross amount"
    tests = wb["7 Test questions"]
    tests["A6"], tests["B6"], tests["C6"], tests["D6"] = "Net sales in May 2026?", "$170,000", "Monthly pack", "May 2026"


def test_the_workbook_holds_what_was_learned_in_the_clients_words(page):
    client, account, _store, _ = page
    wb = _download(client, account)
    assert [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"] == [
        "Start here", "1 Rows to leave out", "2 Metrics", "3 Tables and dates", "4 Links", "5 Codes", "6 Your words",
        "7 Test questions"]
    meta = {r[0]: r[1] for r in wb["_solviz"].iter_rows(values_only=True)}
    assert meta["workspace"] == account
    net = _row(wb["2 Metrics"], "Metric", "Net amount")
    assert net["Add up as"].value == "Total" and net["Shown as"].value == "Money"
    assert str(net[f"Solviz's figure, {_month(wb)}"].value).startswith("$"), "a figure from the data to judge it by"
    assert net["Counted by date"].value == "Order date"
    assert _row(wb["1 Rows to leave out"], "These rows", 'Order lines with status code "C"')["How many"].value
    assert _row(wb["4 Links"], "From", "Return")["Rows without a match"].value == "Kept as Unknown"
    # The ids that bring each answer back are in a hidden column, never in the client's way.
    metrics = wb["2 Metrics"]
    last = metrics.cell(row=4, column=metrics.max_column)
    assert last.value == "Solviz ID (keep)" and metrics.column_dimensions[last.column_letter].hidden


def test_the_answers_are_checked_on_the_data_applied_and_change_the_next_answer(page):
    client, account, store, warehouse = page
    before = _figure(_model(account), warehouse, "Net amount")
    wb = _download(client, account)
    _answer(wb)
    sent = _upload(client, account, wb)
    assert sent.status_code == 303, sent.text[:300]
    check_id = int(sent.headers["location"].rsplit("/", 1)[1])
    items = _items(store, account, check_id)
    leave = next(i for i in items if i["tab"] == "Rows to leave out" and i["ops"])
    assert leave["ok"] and leave["check"].startswith("Leaves out 1,"), leave
    review = client.get(sent.headers["location"])
    assert review.status_code == 200 and "Apply the ticked changes" in review.text and "takings" in review.text
    nothing_yet = _model(account)
    assert not any(f.values == ["C"] for t in nothing_yet.tables.values() for f in t.default_filters)

    chosen = [str(i["n"]) for i in items if i["ops"] and i["ok"]]
    applied = client.post(f"{sent.headers['location']}/apply", data={"item": chosen}, follow_redirects=False)
    assert applied.status_code == 303 and "saved=applied" in applied.headers["location"]

    model = _model(account)
    lines = next(t for t in model.tables.values() if t.key.endswith(".order_lines"))
    assert any(f.op == "ne" and f.values == ["C"] for f in lines.default_filters)
    assert model.columns[f"{lines.key}.status_code"].value_names.get("C") == "Cancelled"
    assert "takings" in next(m for m in model.measures.values() if m.business_name == "Net amount").synonyms["en"]
    assert "turnover" in next(m for m in model.measures.values() if m.business_name == "Gross amount").synonyms["en"]
    link = next(j for j in model.joins.values() if j.from_table.endswith(".returns") and j.to_table.endswith(".customers"))
    assert link.keep_unmatched is False and link.trust == "admin"
    assert _figure(model, warehouse, "Net amount") < before, "cancelled lines no longer count"
    decisions = store.list_core2_overrides(account, None)
    assert decisions and all(d["note"].startswith(f"Setup check #{check_id}") for d in decisions)

    requests = store.list_core2_requests(account, status="waiting")
    assert any(r["source"] == "setup_check" and "freight" in r["note"] for r in requests)
    tests = store.list_core2_test_questions(account, None)
    assert [(t["question"], t["expected"], t["setup_check_id"]) for t in tests] == [
        ("Net sales in May 2026?", "$170,000", check_id)]

    again = client.post(f"{sent.headers['location']}/apply", data={"item": chosen}, follow_redirects=False)
    assert "already" in again.headers["location"], "a setup check is applied once"
    assert len(store.list_core2_test_questions(account, None)) == 1


def test_a_change_the_data_cannot_take_is_shown_as_a_problem_and_never_applied(page):
    client, account, store, _ = page
    wb = _download(client, account)
    net = _row(wb["2 Metrics"], "Metric", "Net amount")
    net["Counted by date"].value = "Delivery date"          # typed over the list: not a date of order lines
    sent = _upload(client, account, wb)
    check_id = int(sent.headers["location"].rsplit("/", 1)[1])
    bad = next(i for i in _items(store, account, check_id) if i["tab"] == "Metrics")
    assert not bad["ok"] and "is not a date of" in bad["check"]
    assert 'class="sc-row sc-bad"' in client.get(sent.headers["location"]).text
    client.post(f"{sent.headers['location']}/apply", data={"item": [str(bad["n"])]}, follow_redirects=False)
    assert not store.list_core2_overrides(account, None)


def test_the_teams_own_figure_is_compared_with_solviz(page):
    client, account, store, _ = page
    wb = _download(client, account)
    net = _row(wb["2 Metrics"], "Metric", "Net amount")
    gross = _row(wb["2 Metrics"], "Metric", "Gross amount")
    month = _month(wb)
    net[f"Your own figure, {month} (optional)"].value = net[f"Solviz's figure, {month}"].value
    gross[f"Your own figure, {month} (optional)"].value = "$1"
    sent = _upload(client, account, wb)
    items = _items(store, account, int(sent.headers["location"].rsplit("/", 1)[1]))
    same = next(i for i in items if i["what"] == "Net amount")
    differs = next(i for i in items if i["what"] == "Gross amount")
    assert same["ok"] and "matches" in same["check"] and not same["ops"]
    assert not differs["ok"] and differs["ops"][0]["op"] == "request"


def test_a_setup_check_from_another_workspace_or_no_setup_check_is_refused(page):
    client, account, _store, _ = page
    wb = _download(client, account)
    wb["_solviz"]["B2"] = "someone-else"
    refused = _upload(client, account, wb)
    assert refused.status_code == 303 and "another%20workspace" in refused.headers["location"]
    from openpyxl import Workbook

    plain = _upload(client, account, Workbook())
    assert "not%20a%20Solviz%20setup%20check" in plain.headers["location"]
    page_text = client.get(plain.headers["location"]).text
    assert "Download the setup check" in page_text and "not a Solviz setup check" in page_text


def test_a_name_that_looks_like_a_formula_is_text_in_the_workbook(page):
    client, account, store, _ = page
    net = next(m for m in _model(account).measures.values() if m.business_name == "Net amount")
    store.set_core2_override(account, None, f"measure:{net.key}", "business_name", '=HYPERLINK("http://x")')
    cell = _row(_download(client, account)["2 Metrics"], "Metric", "=HYPERLINK")["Metric"]
    assert cell.data_type == "s", "a name from the data never runs as a formula in the client's Excel"


def test_a_table_with_no_date_of_its_own_is_counted_by_its_documents_date(retail):
    """A return line with no date of its own is counted by the order line it belongs to, as answers count it: the
    workbook says so, and never asks the client which period its rows belong to."""
    from core2 import setup_check

    _built, learned = retail
    model = learned.model_copy(deep=True)
    returns = next(k for k in model.tables if k.endswith(".returns"))
    for key in [k for k, r in model.date_roles.items() if r.table == returns]:
        del model.date_roles[key]
    model.tables[returns].default_date = None
    assert setup_check._borrowed(model, returns) == "Order date (from Order line)"
    on_returns = [r for r in setup_check.metric_rows(model) if r["key"].startswith(f"m:{returns}")]
    assert on_returns and all(r["date"] == "Order date (from Order line)" and not any(
        "No date found" in n for n in r["notes"]) for r in on_returns)


@pytest.mark.parametrize("name, asked", [
    ("Loan term", True), ("Days past due", True), ("Interest rate", True), ("Credit score", True),
    ("Administration fee", False), ("Days supply", False), ("Net amount", False), ("Percentage of total", True)])
def test_only_a_figure_that_reads_as_a_rate_or_a_score_is_asked_whether_to_average(name, asked):
    from core2.setup_check import _SUMMED_LIKE_A_RATE

    assert bool(_SUMMED_LIKE_A_RATE.search(name)) is asked
