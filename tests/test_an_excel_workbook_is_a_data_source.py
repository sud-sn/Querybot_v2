"""An Excel workbook is a data source: each sheet a table, uploaded in admin or read from a SharePoint link.

Some clients keep their data in workbooks on a shared drive, and a workspace can be tried with no database at all.
A workbook uploaded on the Databases page, or linked from SharePoint or OneDrive (read through Microsoft Graph),
is loaded into DuckDB on this server: each sheet a table, its header the first row that names columns, a title
row above it and a totals row under it left out, each column's type read from its cells. Discovery, Learn and
questions then read it as any other database, in DuckDB's SQL. A linked file is read again when its version
changes, and a link that cannot be read keeps the last copy.

Invented workbooks; Microsoft Graph answered by a stand-in.
"""

from __future__ import annotations

import datetime as dt
import io
import json

import httpx
import pytest
from openpyxl import Workbook

from core import workbook_source as W

ORDERS = [(1001, dt.datetime(2025, 1, 5), "C1", "Widget", 3, 9.5, "Shipped"),
          (1002, dt.datetime(2025, 1, 9), "C2", "Gadget", 1, 120.0, "Shipped"),
          (1003, dt.datetime(2025, 2, 2), "C1", "Widget", 10, 9.5, "Cancelled"),
          (1004, dt.datetime(2025, 2, 20), "C3", "Gizmo", 2, "n/a", "Shipped"),
          (1005, dt.datetime(2025, 3, 1), "C2", "Widget", 4, 9.5, "Shipped")]
CUSTOMERS = [("C1", "Harbor Trading", "North", dt.date(2023, 4, 1)),
             ("C2", "Pine Ridge Foods", "South", dt.date(2024, 6, 1)),
             ("C3", "Quarry Supply Co", "North", dt.date(2024, 9, 9))]


def workbook(orders=ORDERS, *, extra_sheet: bool = False) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Orders"
    sheet.append(["Orders report 2025"])
    sheet.append([])
    sheet.append(["Order ID", "Order Date", "Customer ID", "Product", "Qty", "Unit Price", "Status"])
    for row in orders:
        sheet.append(list(row))
    sheet.append([])
    sheet.append(["Total", None, None, None, sum(r[4] for r in orders), None, None])
    customers = book.create_sheet("Customers")
    customers.append(["Customer ID", "Customer Name", "Region", "Customer Since"])
    for row in CUSTOMERS:
        customers.append(list(row))
    hidden = book.create_sheet("Lookup")
    hidden.append(["a", "b"])
    hidden.append([1, 2])
    hidden.sheet_state = "hidden"
    book.create_sheet("Empty")
    notes = book.create_sheet("Notes")
    notes.append(["Figures are before tax."])
    if extra_sheet:
        regions = book.create_sheet("Regions")
        regions.append(["Region", "Manager Team"])
        regions.append(["North", "Team A"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.fixture()
def store(tmp_path, monkeypatch):
    monkeypatch.setenv("QUERYBOT_WORKBOOKS_DIR", str(tmp_path / "workbooks"))
    W._LINKS.clear()
    return tmp_path / "workbooks"


def test_each_sheet_becomes_a_table(store):
    loaded = W.save(workbook(), "sales.xlsx")
    tables = {t.name: t for t in loaded.tables}
    assert set(tables) == {"orders", "customers"}
    orders = tables["orders"]
    assert [(c.name, c.type) for c in orders.columns] == [
        ("order_id", "BIGINT"), ("order_date", "DATE"), ("customer_id", "VARCHAR"), ("product", "VARCHAR"),
        ("qty", "BIGINT"), ("unit_price", "DOUBLE"), ("status", "VARCHAR")]
    assert orders.columns[1].header == "Order Date" and orders.rows == len(ORDERS) and orders.first_row == 3
    assert orders.left_out == ["2 rows above the header", "a totals row (row 10)"]
    assert [(c.name, c.type) for c in tables["customers"].columns][-1] == ("customer_since", "DATE")
    assert loaded.skipped == ['"Lookup": a hidden sheet', '"Empty": an empty sheet',
                              '"Notes": no header row with rows under it']


def test_figures_read_from_it_are_the_sheets_own(store):
    loaded = W.save(workbook(), "sales.xlsx")
    con = W.connect({"source": W.UPLOAD, "workbook": loaded.sha256})
    shipped = sum(r[4] * r[5] for r in ORDERS if r[6] == "Shipped" and isinstance(r[5], float))
    assert con.execute("SELECT SUM(qty * unit_price) FROM workbook.main.orders WHERE status = 'Shipped'"
                       ).fetchone()[0] == pytest.approx(shipped)
    assert con.execute("SELECT COUNT(*) FROM workbook.main.orders WHERE unit_price IS NULL").fetchone()[0] == 1
    joined = con.execute("SELECT c.region, SUM(o.qty) FROM workbook.main.orders o JOIN workbook.main.customers c "
                         "USING (customer_id) GROUP BY 1 ORDER BY 1").fetchall()
    assert joined == [("North", 15), ("South", 5)]
    comments = dict(con.execute("SELECT column_name, comment FROM duckdb_columns() WHERE table_name = 'orders'"
                                ).fetchall())
    assert comments["order_date"] == "Order Date"
    with pytest.raises(Exception):
        con.execute("DELETE FROM workbook.main.orders")


def test_a_new_version_is_a_new_copy_and_the_same_file_is_loaded_once(store):
    first = W.save(workbook(), "sales.xlsx")
    again = W.save(workbook(), "sales.xlsx")
    assert again.sha256 == first.sha256 and again.loaded_at == first.loaded_at
    newer = W.save(workbook(ORDERS[:3], extra_sheet=True), "sales.xlsx")
    assert newer.sha256 != first.sha256 and newer.shape() != first.shape()
    assert {t.name for t in newer.tables} == {"orders", "customers", "regions"}


@pytest.mark.parametrize("data, says", [(b"order_id,qty\n1,2\n", "not an Excel workbook"),
                                        (b"PK\x03\x04broken", "not an Excel workbook")])
def test_a_file_that_is_no_workbook_is_said_so(store, data, says):
    with pytest.raises(W.WorkbookError, match=says):
        W.save(data, "sales.csv")


def test_names_sql_reads_unquoted():
    assert W.sql_name("Order Date", "x") == "order_date"
    assert W.sql_name("2025 Sales ($)", "x") == "c_2025_sales"
    assert W.sql_name("Coût unitaire", "x") == "cout_unitaire"
    assert W.sql_name("   ", "column_3") == "column_3"


# ── A SharePoint or OneDrive link ────────────────────────────────────────────

LINK = "https://contoso.sharepoint.com/:x:/s/finance/EaBcD123?e=xyz"
CREDS = {"source": W.LINK, "link": LINK, "tenant_id": "t-1", "client_id": "app-1", "client_secret": "s3cret"}


class Graph:
    """Microsoft Graph as the link's file answers: its version tag and its bytes, changed when told."""

    def __init__(self, data: bytes, tag: str = '"{v1},1"'):
        self.data, self.tag, self.calls = data, tag, []
        self.status = 200

    def __call__(self, request: httpx.Request) -> httpx.Response:
        self.calls.append((request.method, request.url.path))
        if "login.microsoftonline.com" in str(request.url):
            return httpx.Response(200, json={"access_token": "tok", "expires_in": 3600})
        assert request.headers["Authorization"] == "Bearer tok"
        assert request.url.path.startswith("/v1.0/shares/u!") and "/driveItem" in request.url.path
        if self.status != 200:
            return httpx.Response(self.status, json={"error": {"message": "Access denied"}})
        if request.url.path.endswith("/content"):
            return httpx.Response(200, content=self.data)
        return httpx.Response(200, json={"name": "Sales.xlsx", "cTag": self.tag, "size": len(self.data),
                                         "lastModifiedDateTime": "2026-10-01T09:00:00Z", "file": {}})

    def downloads(self) -> int:
        return sum(1 for _, path in self.calls if path.endswith("/content"))


def test_a_linked_workbook_is_read_through_graph_and_again_only_when_it_changes(store, monkeypatch):
    graph = Graph(workbook())
    http = httpx.Client(transport=httpx.MockTransport(graph))
    first = W.read_link(CREDS, http=http)
    assert {t.name for t in first.tables} == {"orders", "customers"} and graph.downloads() == 1
    W.read_link(CREDS, http=http)
    assert len([c for c in graph.calls if "graph.microsoft.com" in c[1] or c[1].startswith("/v1.0")]) == 2, \
        "inside the refresh window the copy is used as it is"
    monkeypatch.setattr(W, "refresh_seconds", lambda: 0)
    W.read_link(CREDS, http=http)
    assert graph.downloads() == 1, "the same version tag: not downloaded again"
    graph.data, graph.tag = workbook(ORDERS[:2]), '"{v2},2"'
    newer = W.read_link(CREDS, http=http)
    assert graph.downloads() == 2 and newer.sha256 != first.sha256 and newer.tables[0].rows == 2


def test_a_link_that_cannot_be_read_keeps_the_last_copy_and_says_why_when_none(store, monkeypatch):
    graph = Graph(workbook())
    http = httpx.Client(transport=httpx.MockTransport(graph))
    kept = W.read_link(CREDS, http=http)
    monkeypatch.setattr(W, "refresh_seconds", lambda: 0)
    graph.status = 403
    assert W.read_link(CREDS, http=http).sha256 == kept.sha256
    W._LINKS.clear()
    other = {**CREDS, "link": LINK + "&other=1"}
    with pytest.raises(W.WorkbookError, match="Files.Read.All"):
        W.read_link(other, http=http)


def test_a_link_needs_an_app_to_read_it(store):
    with pytest.raises(W.WorkbookError, match="tenant ID"):
        W.read_link({"source": W.LINK, "link": LINK}, http=httpx.Client(transport=httpx.MockTransport(Graph(b""))))
    with pytest.raises(W.WorkbookError, match="https://"):
        W.read_link({**CREDS, "link": "contoso/sales.xlsx"})


def test_the_share_id_is_graphs_own_encoding():
    # https://learn.microsoft.com/graph/api/shares-get: "u!" + base64url of the link, without its padding.
    assert W.share_id("https://onedrive.live.com/redir?resid=1231244193912!12") == \
        "u!aHR0cHM6Ly9vbmVkcml2ZS5saXZlLmNvbS9yZWRpcj9yZXNpZD0xMjMxMjQ0MTkzOTEyITEy"
    assert json.dumps(W.share_id(LINK)).count("=") == 0
